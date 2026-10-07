import logging
from django.db import transaction
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from .permissions import IsSupervisorOrCommitteeMember
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser

from .models import (
    CustomUser,
    Student,
    Supervisor,
    SupervisorQuota,
    ProjectTopicArea,
    InternshipInfo,
    GraduationProject,
    OutlineReviewGroup,
    OutlineReview,
    WeeklyProgressReport,
    SupervisionMeetingLog,
    SupervisionTask,
    DefenseCouncil,
    CouncilMember,
    CouncilLiveScore,
    EvaluationPolicy,
    FinalGradeSummary,
    AcademicBatch,
    AuditLog,
    Notification,
    Project,
    Document,
    DocumentRequirement,
    Group,
    GroupMember,
    SupervisorOfStudentGroup,
    ProposedAllocation,
    ThesisDeferralRequest,
    DegreeProgram,
)
from .services import (
    NotificationService,
    CouncilConflictService,
    DegreeEligibilityService,
    ThesisAllocationService,
    AcademicClearanceService,
    CouncilStructureService,
)
from .serializers.utc_graduation_serializers import (
    ProjectTopicAreaSerializer,
    SupervisorBriefSerializer,
    InternshipInfoSerializer,
    OutlineReviewSerializer,
    WeeklyProgressReportSerializer,
    SupervisionMeetingLogSerializer,
    SupervisionTaskSerializer,
    CouncilLiveScoreSerializer,
    GraduationProjectDetailSerializer,
    FinalGradeSummarySerializer,
    ProposedAllocationSerializer,
    ThesisDeferralRequestSerializer,
    AcademicBatchSerializer,
)
from .validators import validate_uploaded_file
from .concurrency import retry_on_db_lock

logger = logging.getLogger(__name__)


# ==============================================================================
# MODULE 1: BATCH INITIALIZATION & STUDENT IMPORT APIS (TC_031, TC_032, TC_033)
# ==============================================================================

class BatchCreateAPIView(APIView):
    """
    Khoa khởi tạo đợt đồ án mới, thiết lập thời gian bắt đầu và kết thúc (TC_031).
    Endpoint: /app/batch/create/
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        batches = AcademicBatch.objects.all().order_by("-created_at")
        serializer = AcademicBatchSerializer(batches, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def post(self, request):
        user = request.user
        data = request.data
        batch_code = str(data.get("batch_code", "")).strip()
        batch_name = str(data.get("batch_name", "")).strip()
        start_date = data.get("start_date")
        end_date = data.get("end_date")
        is_active = data.get("is_active", True)
        if isinstance(is_active, str):
            is_active = is_active.lower() in ("true", "1", "yes")

        errors = {}
        if not batch_code:
            errors["batch_code"] = ["Mã đợt đồ án không được để trống."]
        elif AcademicBatch.objects.filter(batch_code=batch_code).exists():
            errors["batch_code"] = ["Mã đợt đồ án đã tồn tại. Vui lòng chọn mã khác."]

        if not batch_name:
            errors["batch_name"] = ["Tên đợt đồ án không được để trống."]

        if start_date and end_date:
            try:
                import datetime
                if isinstance(start_date, str):
                    sd = datetime.date.fromisoformat(start_date)
                else:
                    sd = start_date
                if isinstance(end_date, str):
                    ed = datetime.date.fromisoformat(end_date)
                else:
                    ed = end_date
                if sd > ed:
                    errors["end_date"] = ["Thời gian kết thúc phải sau hoặc bằng thời gian bắt đầu."]
            except (ValueError, TypeError):
                errors["dates"] = ["Định dạng ngày tháng không hợp lệ (YYYY-MM-DD)."]

        if errors:
            return Response(errors, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            batch = AcademicBatch.objects.create(
                batch_code=batch_code,
                batch_name=batch_name,
                start_date=start_date or None,
                end_date=end_date or None,
                is_active=is_active
            )
            # Ensure evaluation policy exists for batch
            EvaluationPolicy.objects.get_or_create(
                batch=batch,
                defaults={
                    "weight_supervisor": 0.4,
                    "weight_reviewer": 0.2,
                    "weight_council": 0.4,
                }
            )

        AuditLog.objects.create(
            user=user,
            action_type="status_change",
            description=f"Khoa khởi tạo đợt đồ án mới: {batch.batch_name} ({batch.batch_code})."
        )

        return Response({
            "success": True,
            "message": f"Khoa khởi tạo đợt đồ án '{batch.batch_name}' thành công.",
            "batch": AcademicBatchSerializer(batch).data
        }, status=status.HTTP_201_CREATED)


class StudentImportAPIView(APIView):
    """
    Khoa Import danh sách sinh viên Cử nhân / Kỹ sư bằng file định dạng mẫu (Excel/CSV).
    - TC_032 (Positive): Import danh sách sinh viên Cử nhân/Kỹ sư thành công từ file mẫu Excel (.xlsx, .xls) hoặc CSV (.csv).
    - TC_033 (Negative): Bắt lỗi và từ chối khi file sai định dạng, thiếu cột bắt buộc, hoặc sai chuẩn dữ liệu (trả về HTTP 400).
    Endpoint: /app/students/import/
    """
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        file_obj = (
            request.FILES.get("file")
            or request.FILES.get("student_file")
            or request.FILES.get("excel_file")
            or request.FILES.get("csv_file")
        )
        if not file_obj:
            return Response({
                "success": False,
                "error": "missing_file",
                "detail": "Vui lòng chọn tệp danh sách sinh viên (.xlsx, .xls hoặc .csv) để tải lên."
            }, status=status.HTTP_400_BAD_REQUEST)

        # 1. Check file extension
        import os
        ext = os.path.splitext(file_obj.name)[1].lower()
        if ext not in [".xlsx", ".xls", ".csv"]:
            return Response({
                "success": False,
                "error": "invalid_file_format",
                "detail": f"Định dạng tệp '{ext}' không được hỗ trợ. Hệ thống chỉ chấp nhận tệp Excel (.xlsx, .xls) hoặc CSV (.csv)."
            }, status=status.HTTP_400_BAD_REQUEST)

        # 2. Identify target AcademicBatch
        batch_id = request.data.get("batch_id")
        batch = None
        if batch_id:
            batch = AcademicBatch.objects.filter(id=batch_id).first()
        if not batch:
            batch = AcademicBatch.objects.filter(is_active=True).first() or AcademicBatch.objects.last()

        # 3. Parse content
        rows = []
        raw_headers = []
        try:
            if ext in [".xlsx", ".xls"]:
                import openpyxl
                file_obj.seek(0)
                wb = openpyxl.load_workbook(file_obj, data_only=True)
                sheet = wb.active
                all_rows = list(sheet.iter_rows(values_only=True))
                if not all_rows or len(all_rows) < 2:
                    return Response({
                        "success": False,
                        "error": "empty_file",
                        "detail": "Tệp Excel không chứa dữ liệu hoặc chỉ có tiêu đề."
                    }, status=status.HTTP_400_BAD_REQUEST)
                raw_headers = [str(cell).strip() if cell is not None else "" for cell in all_rows[0]]
                for r in all_rows[1:]:
                    if any(cell is not None and str(cell).strip() != "" for cell in r):
                        row_dict = {}
                        for h_idx, h in enumerate(raw_headers):
                            if h:
                                val = r[h_idx] if h_idx < len(r) else None
                                row_dict[h] = val
                        rows.append(row_dict)
            else:  # .csv
                import csv
                import io
                file_obj.seek(0)
                content = file_obj.read()
                for enc in ["utf-8-sig", "utf-8", "latin-1"]:
                    try:
                        text_stream = io.StringIO(content.decode(enc))
                        reader = csv.DictReader(text_stream)
                        raw_headers = reader.fieldnames or []
                        rows = [r for r in reader if any(v and v.strip() for v in r.values() if v is not None)]
                        break
                    except UnicodeDecodeError:
                        continue

                if not rows:
                    return Response({
                        "success": False,
                        "error": "empty_file",
                        "detail": "Tệp CSV không chứa dữ liệu hoặc chỉ có tiêu đề."
                    }, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response({
                "success": False,
                "error": "parse_error",
                "detail": f"Không thể đọc nội dung tệp: {str(e)}"
            }, status=status.HTTP_400_BAD_REQUEST)

        # 4. Normalize and validate required columns (TC_033)
        def normalize_col_name(c):
            import unicodedata
            clean = str(c).replace("đ", "d").replace("Đ", "D")
            norm = unicodedata.normalize('NFKD', clean).encode('ASCII', 'ignore').decode('utf-8')
            return norm.lower().strip().replace(" ", "_").replace(".", "").replace("-", "_")

        header_norm_map = {normalize_col_name(h): h for h in raw_headers if h}

        # Match columns
        reg_col = None
        for candidate in ["registration_no", "student_id", "ma_sv", "mssv", "ma_sinh_vien", "masv"]:
            if candidate in header_norm_map:
                reg_col = header_norm_map[candidate]
                break

        name_col = None
        for candidate in ["full_name", "ho_ten", "ho_va_ten", "ten_sinh_vien", "ho_ten_sinh_vien", "name", "ten"]:
            if candidate in header_norm_map:
                name_col = header_norm_map[candidate]
                break

        prog_col = None
        for candidate in ["degree_program", "he_dao_tao", "he", "chuong_trinh", "chuong_trinh_dao_tao", "program", "program_type"]:
            if candidate in header_norm_map:
                prog_col = header_norm_map[candidate]
                break

        email_col = None
        for candidate in ["email", "thu_dien_tu", "email_lms"]:
            if candidate in header_norm_map:
                email_col = header_norm_map[candidate]
                break

        class_col = None
        for candidate in ["class_name", "lop", "lop_hoc_phan", "lop_sinh_hoat", "department", "khoa"]:
            if candidate in header_norm_map:
                class_col = header_norm_map[candidate]
                break

        missing_columns = []
        if not reg_col:
            missing_columns.append("Mã sinh viên (registration_no / ma_sv / mssv)")
        if not name_col:
            missing_columns.append("Họ và tên (full_name / ho_ten)")
        if not prog_col:
            missing_columns.append("Hệ đào tạo (degree_program / he_dao_tao)")

        if missing_columns:
            return Response({
                "success": False,
                "error": "missing_required_columns",
                "detail": f"Tệp danh sách thiếu các cột bắt buộc: {', '.join(missing_columns)}. Vui lòng sử dụng file mẫu chuẩn.",
                "missing_columns": missing_columns,
                "found_columns": raw_headers
            }, status=status.HTTP_400_BAD_REQUEST)

        # 5. Row-level data validation (TC_033)
        row_errors = []
        valid_records = []

        for idx, row in enumerate(rows, start=2):
            reg_val = str(row.get(reg_col) or "").strip()
            name_val = str(row.get(name_col) or "").strip()
            prog_val = str(row.get(prog_col) or "").strip()
            email_val = str(row.get(email_col) or "").strip() if email_col else ""
            dept_val = str(row.get(class_col) or "").strip() if class_col else "CNTT"

            if not reg_val:
                row_errors.append(f"Dòng {idx}: Mã sinh viên không được để trống.")
                continue

            if not name_val:
                row_errors.append(f"Dòng {idx}: Họ và tên sinh viên không được để trống.")
                continue

            # Standardize degree program
            import unicodedata
            clean_prog = str(prog_val).replace("đ", "d").replace("Đ", "D")
            norm_prog = unicodedata.normalize('NFKD', clean_prog).encode('ASCII', 'ignore').decode('utf-8').upper()
            if any(k in norm_prog for k in ["KY SU", "ENGINEER", "KS"]):
                degree_program = DegreeProgram.ENGINEER
            elif any(k in norm_prog for k in ["CU NHAN", "BACHELOR", "CN"]):
                degree_program = DegreeProgram.BACHELOR
            else:
                row_errors.append(f"Dòng {idx}: Hệ đào tạo '{prog_val}' không hợp lệ (Chỉ chấp nhận 'Kỹ sư' hoặc 'Cử nhân').")
                continue

            # Check email syntax if provided
            if email_val:
                if "@" not in email_val or "." not in email_val.split("@")[-1]:
                    row_errors.append(f"Dòng {idx}: Email '{email_val}' không đúng định dạng.")
                    continue
            else:
                email_val = f"{reg_val.lower()}@lms.utc.edu.vn"

            valid_records.append({
                "row_number": idx,
                "registration_no": reg_val,
                "full_name": name_val,
                "degree_program": degree_program,
                "email": email_val,
                "department": dept_val,
            })

        # If there are data validation errors (TC_033), reject the import
        if row_errors:
            return Response({
                "success": False,
                "error": "data_validation_failed",
                "detail": f"Dữ liệu trong tệp không đạt chuẩn ({len(row_errors)} lỗi phát hiện). Vui lòng kiểm tra lại dữ liệu.",
                "errors_count": len(row_errors),
                "errors": row_errors[:20]
            }, status=status.HTTP_400_BAD_REQUEST)

        # 6. Execute Import (TC_032 - Positive)
        imported_students = []
        engineer_count = 0
        bachelor_count = 0

        with transaction.atomic():
            for rec in valid_records:
                reg_no = rec["registration_no"]
                full_name = rec["full_name"]
                email = rec["email"]
                deg_prog = rec["degree_program"]
                dept = rec["department"]

                parts = full_name.split()
                if len(parts) > 1:
                    first_name = parts[-1]
                    last_name = " ".join(parts[:-1])
                else:
                    first_name = full_name
                    last_name = ""

                user, u_created = CustomUser.objects.update_or_create(
                    username=reg_no,
                    defaults={
                        "email": email,
                        "first_name": first_name,
                        "last_name": last_name,
                        "user_type": "student",
                    }
                )
                if u_created:
                    user.set_password("Utc@123456")
                    user.save()

                student, s_created = Student.objects.update_or_create(
                    user=user,
                    defaults={
                        "registration_no": reg_no,
                        "degree_program": deg_prog,
                        "department": dept,
                        "academic_batch": batch,
                    }
                )

                if deg_prog == DegreeProgram.ENGINEER:
                    engineer_count += 1
                else:
                    bachelor_count += 1

                imported_students.append({
                    "registration_no": student.registration_no,
                    "full_name": user.get_full_name() or full_name,
                    "degree_program": student.degree_program,
                    "degree_program_display": student.get_degree_program_display(),
                    "email": user.email,
                    "department": student.department,
                    "batch": batch.batch_code if batch else None,
                })

        AuditLog.objects.create(
            user=request.user,
            action_type="status_change",
            description=f"Khoa import thành công {len(imported_students)} sinh viên (Kỹ sư: {engineer_count}, Cử nhân: {bachelor_count})."
        )

        return Response({
            "success": True,
            "message": f"Khoa đã import thành công {len(imported_students)} sinh viên (Kỹ sư: {engineer_count}, Cử nhân: {bachelor_count}).",
            "total_imported": len(imported_students),
            "engineer_count": engineer_count,
            "bachelor_count": bachelor_count,
            "batch_code": batch.batch_code if batch else None,
            "students": imported_students
        }, status=status.HTTP_201_CREATED)


# ==============================================================================
# STUDENT SURVEY & ONBOARDING API
# ==============================================================================

class StudentSurveyAPIView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get(self, request):
        user = request.user
        if user.user_type != "student":
            return Response({"detail": "Chỉ dành cho tài khoản sinh viên."}, status=status.HTTP_403_FORBIDDEN)

        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        batch = student.academic_batch or AcademicBatch.objects.filter(is_active=True).first()
        
        from django.core.cache import cache

        topic_areas_data = cache.get("utc_active_topic_areas_data")
        if topic_areas_data is None:
            topic_areas = ProjectTopicArea.objects.filter(is_active=True)
            topic_areas_data = ProjectTopicAreaSerializer(topic_areas, many=True).data
            cache.set("utc_active_topic_areas_data", topic_areas_data, 600)

        supervisors_data = cache.get("utc_active_supervisors_brief_data")
        if supervisors_data is None:
            supervisors = Supervisor.objects.all().select_related("user").order_by("user__first_name")
            supervisors_data = SupervisorBriefSerializer(supervisors, many=True).data
            cache.set("utc_active_supervisors_brief_data", supervisors_data, 300)

        survey = InternshipInfo.objects.filter(student=student).first()
        survey_data = InternshipInfoSerializer(survey).data if survey else None

        return Response({
            "student": {
                "id": student.id,
                "registration_no": student.registration_no,
                "full_name": user.get_full_name() or user.username,
                "email": user.email,
                "phone_number": student.phone_number,
                "department": student.department,
                "degree_program": student.degree_program,
                "degree_program_display": student.get_degree_program_display(),
                "cpa": student.cpa,
                "credits_accumulated": student.credits_accumulated,
                "is_eligible_for_thesis": student.is_eligible_for_thesis,
                "course_class": student.course_class.class_name if student.course_class else ""
            },
            "batch": {
                "id": getattr(batch, "id", None) if batch else None,
                "batch_code": getattr(batch, "batch_code", "") if batch else "",
                "batch_name": getattr(batch, "batch_name", "") if batch else ""
            },
            "topic_areas": topic_areas_data,
            "supervisors": supervisors_data,
            "survey": survey_data
        }, status=status.HTTP_200_OK)

    def post(self, request):
        user = request.user
        if user.user_type != "student":
            return Response({"detail": "Chỉ dành cho tài khoản sinh viên."}, status=status.HTTP_403_FORBIDDEN)

        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        batch = student.academic_batch or AcademicBatch.objects.filter(is_active=True).first()
        if not batch:
            return Response({"detail": "Không có đợt làm đồ án nào đang hoạt động."}, status=status.HTTP_400_BAD_REQUEST)

        data = request.data
        is_interning = str(data.get("is_interning", "false")).lower() in ["true", "1"]
        company_name = data.get("company_name", "").strip()
        topic_direction_id = data.get("topic_direction")
        
        pref1_id = data.get("preference_1") or data.get("preferred_supervisor")
        pref2_id = data.get("preference_2")
        pref3_id = data.get("preference_3")
        secondary_criteria_note = data.get("secondary_criteria_note", "").strip()

        tentative_title = data.get("tentative_title", "").strip()
        phone_number = data.get("phone_number", "").strip()
        email = data.get("email", "").strip()
        new_password = data.get("new_password", "").strip()

        # Validation
        if is_interning and not company_name:
            return Response({"company_name": ["Vui lòng nhập tên công ty/doanh nghiệp đang thực tập."]}, status=status.HTTP_400_BAD_REQUEST)

        topic_direction = ProjectTopicArea.objects.filter(id=topic_direction_id).first() if topic_direction_id else None
        
        sup1 = Supervisor.objects.filter(id=pref1_id).first() if pref1_id else None
        sup2 = Supervisor.objects.filter(id=pref2_id).first() if pref2_id else None
        sup3 = Supervisor.objects.filter(id=pref3_id).first() if pref3_id else None

        # Check duplicate preferences
        selected_sups = [s for s in [sup1, sup2, sup3] if s is not None]
        selected_ids = [s.id for s in selected_sups]
        if len(selected_ids) != len(set(selected_ids)):
            return Response({"detail": "Các nguyện vọng NV1, NV2, NV3 không được trùng nhau. Vui lòng chọn các thầy cô khác nhau!"}, status=status.HTTP_400_BAD_REQUEST)

        # Nút quyết định Giai đoạn 2: SV Kỹ sư -> Kiểm tra học vị GV tối thiểu (Tiến sĩ trở lên)
        try:
            DegreeEligibilityService.validate_preferences_for_student(student, selected_sups)
        except ValueError as ve:
            return Response({"detail": str(ve)}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            # Update student profile
            if phone_number:
                student.phone_number = phone_number
                student.save(update_fields=["phone_number"])

            if email:
                user.email = email
            if new_password and len(new_password) >= 6:
                user.set_password(new_password)
            user.save()

            # Update or create InternshipInfo
            survey, _ = InternshipInfo.objects.update_or_create(
                student=student,
                defaults={
                    "batch": batch,
                    "is_interning": is_interning,
                    "company_name": company_name if is_interning else "",
                    "topic_direction": topic_direction,
                    "preferred_supervisor": sup1,
                    "preference_1": sup1,
                    "preference_2": sup2,
                    "preference_3": sup3,
                    "secondary_criteria_note": secondary_criteria_note,
                    "tentative_title": tentative_title
                }
            )

        return Response({
            "message": "Cập nhật khảo sát và nguyện vọng thành công!",
            "survey": InternshipInfoSerializer(survey).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# GRADUATION PROJECT & OUTLINE WORKFLOW
# ==============================================================================

class StudentGraduationProjectAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).select_related(
            "supervisor__user",
            "reviewer__user",
            "council",
            "topic_category",
            "final_grade_summary"
        ).first()

        if not project:
            return Response({
                "has_project": False,
                "message": "Bạn chưa được phân công Đề tài và Giảng viên hướng dẫn."
            }, status=status.HTTP_200_OK)

        return Response({
            "has_project": True,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class StudentOutlineSubmissionAPIView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    @retry_on_db_lock(max_retries=5, initial_delay=0.05, backoff_factor=1.5)
    def post(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            return Response({"detail": "Chưa được phân công đề tài."}, status=status.HTTP_400_BAD_REQUEST)

        topic_title_vi = request.data.get("topic_title_vi", "").strip()
        topic_title_en = request.data.get("topic_title_en", "").strip()
        outline_file = request.FILES.get("outline_file")

        if not topic_title_vi:
            return Response({"topic_title_vi": ["Vui lòng nhập tên đề tài tiếng Việt."]}, status=status.HTTP_400_BAD_REQUEST)

        # Check for duplicated topic names
        from .models import SupervisorOfStudentGroup
        is_duplicated = (
            GraduationProject.objects.filter(topic_title_vi__iexact=topic_title_vi, status="PASSED").exclude(id=project.id).exists() or
            SupervisorOfStudentGroup.objects.filter(project__project_name__iexact=topic_title_vi, status="accepted").exists()
        )
        if is_duplicated:
            return Response({"topic_title_vi": ["Tên đề tài đã trùng lặp với đề tài đã được nghiệm thu từ các năm trước."]}, status=status.HTTP_400_BAD_REQUEST)

        # File validation with binary magic checks
        if outline_file:
            try:
                validate_uploaded_file(outline_file, allowed_extensions=[".pdf"], max_size_bytes=25 * 1024 * 1024)
            except Exception as e:
                err_msg = getattr(e, "detail", str(e))
                if isinstance(err_msg, list):
                    err_msg = err_msg[0]
                return Response({"outline_file": [str(err_msg)]}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            project = GraduationProject.objects.select_for_update().get(id=project.id)
            project.topic_title_vi = topic_title_vi
            if topic_title_en:
                project.topic_title_en = topic_title_en
            project.status = "OUTLINE_PENDING"
            project.save()

            review, _ = OutlineReview.objects.select_for_update().get_or_create(project=project)
            if outline_file:
                review.outline_file = outline_file
            review.verdict = "PENDING"
            review.save()

        return Response({
            "message": "Nộp đề cương thành công, đang chờ Giảng viên xét duyệt!",
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# WEEKLY PROGRESS REPORTS (Week 1 -> 15)
# ==============================================================================

class StudentWeeklyReportAPIView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            return Response({"detail": "Chưa có đồ án."}, status=status.HTTP_400_BAD_REQUEST)

        reports = WeeklyProgressReport.objects.filter(project=project).order_by("week_number")
        return Response(WeeklyProgressReportSerializer(reports, many=True).data, status=status.HTTP_200_OK)

    @retry_on_db_lock(max_retries=5, initial_delay=0.05, backoff_factor=1.5)
    def post(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            return Response({"detail": "Chưa có đồ án."}, status=status.HTTP_400_BAD_REQUEST)

        week_number = request.data.get("week_number")
        summary_content = request.data.get("summary_content", "").strip()
        planned_tasks = request.data.get("planned_tasks", "").strip()
        git_commit_link = request.data.get("git_commit_link", "").strip()
        attached_file = request.FILES.get("attached_file")

        try:
            week_num = int(week_number)
            if week_num < 1 or week_num > 15:
                return Response({"week_number": ["Tuần báo cáo phải từ 1 đến 15."]}, status=status.HTTP_400_BAD_REQUEST)
        except (ValueError, TypeError):
            return Response({"week_number": ["Tuần báo cáo không hợp lệ."]}, status=status.HTTP_400_BAD_REQUEST)

        if not summary_content:
            return Response({"summary_content": ["Vui lòng nhập nội dung tóm tắt kết quả công việc trong tuần."]}, status=status.HTTP_400_BAD_REQUEST)

        # File validation with binary magic checks
        if attached_file:
            try:
                validate_uploaded_file(
                    attached_file,
                    allowed_extensions=[".pdf", ".zip", ".rar", ".docx", ".xlsx", ".pptx"],
                    max_size_bytes=25 * 1024 * 1024
                )
            except Exception as e:
                err_msg = getattr(e, "detail", str(e))
                if isinstance(err_msg, list):
                    err_msg = err_msg[0]
                return Response({"attached_file": [str(err_msg)]}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            project = GraduationProject.objects.select_for_update().get(id=project.id)
            report, _ = WeeklyProgressReport.objects.select_for_update().update_or_create(
                project=project,
                week_number=week_num,
                defaults={
                    "summary_content": summary_content,
                    "planned_tasks": planned_tasks,
                    "git_commit_link": git_commit_link,
                    "submitted_at": timezone.now(),
                }
            )
            if attached_file:
                report.attached_file = attached_file
                report.save(update_fields=["attached_file", "submitted_at"])

        return Response({
            "message": f"Nộp báo cáo tuần {week_num} thành công!",
            "report": WeeklyProgressReportSerializer(report).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# STUDENT SUPERVISION LOGS & TASK BOARD APIS
# ==============================================================================

class StudentSupervisionLogsAPIView(APIView):
    """Sinh viên xem danh sách nhật ký các buổi gặp / làm việc từ GVHD"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            membership = GroupMember.objects.filter(student=student).first()
            if membership and membership.group:
                project = GraduationProject.objects.filter(student__group_memberships__group=membership.group).first()

        if not project:
            return Response([], status=status.HTTP_200_OK)

        logs = SupervisionMeetingLog.objects.filter(project=project).order_by("-meeting_date", "-created_at")
        return Response(SupervisionMeetingLogSerializer(logs, many=True).data, status=status.HTTP_200_OK)



class StudentTasksAPIView(APIView):
    """Sinh viên xem danh sách công việc được giao và tiến độ tổng quan"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            return Response({"detail": "Chưa được phân công đề tài."}, status=status.HTTP_400_BAD_REQUEST)

        tasks = SupervisionTask.objects.filter(project=project).order_by("is_completed", "due_date", "-created_at")
        total = tasks.count()
        completed = tasks.filter(is_completed=True).count()
        in_progress = tasks.filter(status="IN_PROGRESS", is_completed=False).count()
        todo = tasks.filter(status="TODO", is_completed=False).count()
        completion_rate = round((completed / total * 100), 1) if total > 0 else 0

        return Response({
            "stats": {
                "total": total,
                "completed": completed,
                "in_progress": in_progress,
                "todo": todo,
                "completion_rate": completion_rate,
            },
            "tasks": SupervisionTaskSerializer(tasks, many=True).data
        }, status=status.HTTP_200_OK)


class StudentMarkTaskCompletedAPIView(APIView):
    """Sinh viên đánh dấu hoàn thành nhiệm vụ (hoặc bỏ chọn) kèm ghi chú kết quả"""
    permission_classes = [IsAuthenticated]

    def patch(self, request, pk):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Hồ sơ sinh viên không tồn tại."}, status=status.HTTP_404_NOT_FOUND)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            return Response({"detail": "Chưa được phân công đề tài."}, status=status.HTTP_400_BAD_REQUEST)

        task = get_object_or_404(SupervisionTask, id=pk, project=project)

        is_completed_val = request.data.get("is_completed")
        if is_completed_val is not None:
            is_completed = str(is_completed_val).lower() in ["true", "1"]
        else:
            is_completed = not task.is_completed

        student_notes = request.data.get("student_notes")
        if student_notes is not None:
            task.student_notes = str(student_notes).strip()

        task.is_completed = is_completed
        if is_completed:
            task.status = "COMPLETED"
            task.completed_at = timezone.now()
        else:
            task.status = "IN_PROGRESS"
            task.completed_at = None
        task.save()

        try:
            status_text = "đã hoàn thành" if is_completed else "đang thực hiện lại"
            NotificationService.create_notification(
                user=project.supervisor.user,
                notification_type="general",
                title=f"[Tiến độ nhiệm vụ] SV {student.user.get_full_name()}",
                message=f"Sinh viên {student.user.get_full_name()} ({student.registration_no}) {status_text} nhiệm vụ: '{task.title}'.",
            )
        except Exception as e:
            logger.warning("Could not send notification for task completion: %s", e)

        all_tasks = SupervisionTask.objects.filter(project=project)
        total = all_tasks.count()
        completed = all_tasks.filter(is_completed=True).count()
        rate = round((completed / total * 100), 1) if total > 0 else 0

        return Response({
            "message": f"Đã cập nhật trạng thái nhiệm vụ: {'Hoàn thành' if is_completed else 'Chưa hoàn thành'}",
            "task": SupervisionTaskSerializer(task).data,
            "stats": {
                "total": total,
                "completed": completed,
                "completion_rate": rate,
            }
        }, status=status.HTTP_200_OK)


# ==============================================================================
# SUPERVISOR DASHBOARD & ACTIONS
# ==============================================================================

class SupervisorGraduationProjectsAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên hướng dẫn."}, status=status.HTTP_403_FORBIDDEN)

        projects = GraduationProject.objects.filter(supervisor=supervisor).select_related(
            "student__user",
            "topic_category",
            "council",
            "final_grade_summary"
        ).order_by("student__user__last_name")

        return Response(GraduationProjectDetailSerializer(projects, many=True).data, status=status.HTTP_200_OK)

class SupervisorOutlineGroupReviewListAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        # Supervisor can review outlines from:
        # 1. Any review groups they are a member of
        # 2. Any graduation projects they directly supervise
        review_groups = supervisor.outline_groups.all()

        # Ensure outline reviews exist for projects supervised by this supervisor
        supervised_projects = GraduationProject.objects.filter(supervisor=supervisor)
        for p in supervised_projects:
            OutlineReview.objects.get_or_create(project=p)

        from django.db.models import Q
        query = Q(project__supervisor=supervisor)
        if review_groups.exists():
            query |= Q(review_group__in=review_groups)

        reviews = OutlineReview.objects.filter(query).select_related(
            "project__student__user",
            "project__supervisor__user",
            "review_group",
            "reviewer__user"
        ).distinct()

        from .utils.vietnamese_sort import sort_by_vietnamese_name
        reviews_list = list(reviews)
        reviews_list = sort_by_vietnamese_name(
            reviews_list,
            key_extractor=lambda r: (f"{r.project.student.user.last_name} {r.project.student.user.first_name}".strip() if (r.project.student.user.last_name or r.project.student.user.first_name) else (r.project.student.user.get_full_name() or r.project.student.user.username))
        )
        
        from app.serializers.utc_graduation_serializers import OutlineReviewSerializer
        return Response(OutlineReviewSerializer(reviews_list, many=True).data, status=status.HTTP_200_OK)


class SupervisorOutlineReviewAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.data.get("project_id")
        verdict = request.data.get("verdict")  # APPROVED, REVISION_REQUIRED, REJECTED
        comments = request.data.get("comments", "").strip()

        if not project_id or not verdict:
            return Response({"detail": "project_id và verdict là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id)
        
        # Verify permissions: either supervisor of the project OR member of the assigned review group
        is_supervisor = project.supervisor_id == supervisor.id
        review, _ = OutlineReview.objects.get_or_create(project=project)
        
        is_group_reviewer = False
        if review.review_group and review.review_group.members.filter(id=supervisor.id).exists():
            is_group_reviewer = True
            
        if not is_supervisor and not is_group_reviewer:
            return Response({"detail": "Bạn không có quyền thẩm định đề cương này."}, status=status.HTTP_403_FORBIDDEN)

        with transaction.atomic():
            review.reviewer = supervisor
            review.verdict = verdict
            review.comments = comments
            review.reviewed_at = timezone.now()
            review.save()

            if verdict == "APPROVED":
                project.status = "OUTLINE_APPROVED"
            elif verdict == "REVISION_REQUIRED":
                project.status = "OUTLINE_REVISION"
            elif verdict == "REJECTED":
                project.status = "FAILED"
            project.save()

        verdict_display = getattr(review, "get_verdict_display", lambda: verdict)()
        return Response({
            "message": f"Đã cập nhật kết quả duyệt đề cương: {verdict_display}",
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class SupervisorWeeklyFeedbackAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        report_id = request.data.get("report_id")
        rating = request.data.get("rating")  # GOOD, ACCEPTABLE, LATE, UNSATISFACTORY
        feedback = request.data.get("feedback", "").strip()

        if not report_id or not rating:
            return Response({"detail": "report_id và rating là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        report = get_object_or_404(WeeklyProgressReport, id=report_id, project__supervisor=supervisor)
        report.supervisor_rating = rating
        report.supervisor_feedback = feedback
        report.reviewed_at = timezone.now()
        report.save()

        return Response({
            "message": "Đã lưu nhận xét và đánh giá tuần!",
            "report": WeeklyProgressReportSerializer(report).data
        }, status=status.HTTP_200_OK)


class SupervisorDefenseEvaluationAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên hướng dẫn."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.data.get("project_id")
        supervisor_score = request.data.get("supervisor_score")
        supervisor_feedback = request.data.get("supervisor_feedback", "").strip()
        is_eligible = request.data.get("is_eligible_for_defense", True)

        try:
            score = float(supervisor_score)
            if score < 0.0 or score > 10.0:
                return Response({"supervisor_score": ["Điểm hướng dẫn phải từ 0.0 đến 10.0"]}, status=status.HTTP_400_BAD_REQUEST)
        except (ValueError, TypeError):
            return Response({"supervisor_score": ["Điểm số không hợp lệ."]}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)
        is_draft = bool(request.data.get("is_draft", False))

        with transaction.atomic():
            project.supervisor_score = score
            project.supervisor_feedback = supervisor_feedback
            project.supervisor_score_is_draft = is_draft
            if not is_draft:
                project.is_eligible_for_defense = is_eligible
                if is_eligible:
                    project.status = "DEFENSE_READY"
                project.save()

                # Update Final Grade
                summary, _ = FinalGradeSummary.objects.get_or_create(project=project)
                summary.supervisor_score = score
                policy = EvaluationPolicy.objects.filter(batch=project.batch).first()
                summary.calculate_and_save(policy=policy)
            else:
                project.save()

        msg = "Đã lưu nháp phiếu đánh giá của Giảng viên hướng dẫn! Điểm chưa công bố cho sinh viên." if is_draft else "Đã lưu và công bố phiếu đánh giá của Giảng viên hướng dẫn!"
        return Response({
            "message": msg,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class SupervisorSupervisionLogsAPIView(APIView):
    """Giảng viên xem và tạo nhật ký làm việc / họp với sinh viên"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.query_params.get("project_id")
        if project_id:
            project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)
            logs = SupervisionMeetingLog.objects.filter(project=project).order_by("-meeting_date", "-created_at")
        else:
            logs = SupervisionMeetingLog.objects.filter(project__supervisor=supervisor).order_by("-meeting_date", "-created_at")

        return Response(SupervisionMeetingLogSerializer(logs, many=True).data, status=status.HTTP_200_OK)

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.data.get("project_id")
        meeting_date = request.data.get("meeting_date")
        meeting_time = request.data.get("meeting_time", "09:00 - 10:30")
        meeting_type = request.data.get("meeting_type", "OFFLINE")
        location_or_link = request.data.get("location_or_link", "").strip()
        content_discussed = request.data.get("content_discussed", "").strip()
        supervisor_notes = request.data.get("supervisor_notes", "").strip()
        next_meeting_plan = request.data.get("next_meeting_plan", "").strip()

        if not project_id or not meeting_date or not content_discussed:
            return Response({"detail": "project_id, meeting_date và content_discussed là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)

        with transaction.atomic():
            log = SupervisionMeetingLog.objects.create(
                project=project,
                meeting_date=meeting_date,
                meeting_time=meeting_time,
                meeting_type=meeting_type,
                location_or_link=location_or_link,
                content_discussed=content_discussed,
                supervisor_notes=supervisor_notes,
                next_meeting_plan=next_meeting_plan,
            )

            # Feature 5: Giao việc tuần tới kèm Deadline làm căn cứ đánh giá điểm quá trình
            task_title = request.data.get("task_title", "").strip()
            created_task = None
            if task_title:
                task_desc = request.data.get("task_description", "").strip()
                task_due_date = request.data.get("task_due_date") or None
                task_priority = request.data.get("task_priority", "MEDIUM")
                created_task = SupervisionTask.objects.create(
                    project=project,
                    meeting_log=log,
                    title=task_title,
                    description=task_desc,
                    assigned_by=supervisor,
                    due_date=task_due_date,
                    priority=task_priority,
                    status="TODO",
                )

        try:
            supervisor_name = supervisor.user.get_full_name() or supervisor.user.username
            if meeting_type == "ONLINE" and location_or_link:
                notif_title = f"[Lịch họp trực tuyến] Cuộc hẹn từ GVHD {supervisor_name}"
                notif_msg = f"GVHD đã tạo lịch hẹn gặp trực tuyến lúc {meeting_time} ngày {meeting_date}. Link tham gia: {location_or_link}"
            else:
                notif_title = "[Nhật ký hướng dẫn] GVHD vừa ghi nhận buổi làm việc"
                notif_msg = f"GVHD {supervisor_name} đã cập nhật nhật ký buổi làm việc ngày {meeting_date}."

            NotificationService.create_notification(
                user=project.student.user,
                notification_type="general",
                title=notif_title,
                message=notif_msg,
                action_url="/student/dashboard",
                send_email=True,
            )

            if created_task:
                NotificationService.create_notification(
                    user=project.student.user,
                    notification_type="general",
                    title=f"[Nhiệm vụ mới] {task_title}",
                    message=f"GVHD đã giao nhiệm vụ tuần tới: '{task_title}'. Hạn nộp: {created_task.due_date or 'Không có'}. (Lưu làm căn cứ đánh giá điểm quá trình)",
                    action_url="/student/dashboard",
                    send_email=True,
                )
        except Exception as e:
            logger.warning("Could not send notification for meeting log: %s", e)

        return Response({
            "message": "Đã lưu nhật ký hướng dẫn và căn cứ đánh giá điểm quá trình thành công!",
            "log": SupervisionMeetingLogSerializer(log).data
        }, status=status.HTTP_201_CREATED)


class SupervisorTasksAPIView(APIView):
    """Giảng viên xem, giao nhiệm vụ mới, sửa hoặc xóa nhiệm vụ cho sinh viên"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.query_params.get("project_id")
        if project_id:
            project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)
            tasks = SupervisionTask.objects.filter(project=project).order_by("is_completed", "due_date", "-created_at")
        else:
            tasks = SupervisionTask.objects.filter(project__supervisor=supervisor).order_by("is_completed", "due_date", "-created_at")

        return Response(SupervisionTaskSerializer(tasks, many=True).data, status=status.HTTP_200_OK)

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.data.get("project_id")
        title = request.data.get("title", "").strip()
        description = request.data.get("description", "").strip()
        due_date = request.data.get("due_date") or None
        priority = request.data.get("priority", "MEDIUM")
        meeting_log_id = request.data.get("meeting_log_id")

        if not project_id or not title:
            return Response({"detail": "project_id và title là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)
        meeting_log = SupervisionMeetingLog.objects.filter(id=meeting_log_id, project=project).first() if meeting_log_id else None

        task = SupervisionTask.objects.create(
            project=project,
            meeting_log=meeting_log,
            title=title,
            description=description,
            assigned_by=supervisor,
            due_date=due_date,
            priority=priority,
            status="TODO",
            is_completed=False,
        )

        try:
            NotificationService.create_notification(
                user=project.student.user,
                notification_type="general",
                title="[Nhiệm vụ mới] GVHD vừa giao việc cho bạn",
                message=f"GVHD {supervisor.user.get_full_name()} đã giao nhiệm vụ: '{title}'. Hạn nộp: {due_date or 'Không'}.",
            )
        except Exception as e:
            logger.warning("Could not send notification for task creation: %s", e)

        return Response({
            "message": f"Đã giao nhiệm vụ '{title}' cho sinh viên!",
            "task": SupervisionTaskSerializer(task).data
        }, status=status.HTTP_201_CREATED)

    def patch(self, request, pk=None):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        task_id = pk or request.data.get("task_id")
        task = get_object_or_404(SupervisionTask, id=task_id, project__supervisor=supervisor)

        for field in ["title", "description", "due_date", "priority", "status", "is_completed"]:
            if field in request.data:
                val = request.data[field]
                if field == "is_completed":
                    val = str(val).lower() in ["true", "1"]
                setattr(task, field, val)

        if task.is_completed and not task.completed_at:
            task.completed_at = timezone.now()
        elif not task.is_completed:
            task.completed_at = None

        task.save()
        return Response({
            "message": "Đã cập nhật nhiệm vụ thành công!",
            "task": SupervisionTaskSerializer(task).data
        }, status=status.HTTP_200_OK)

    def delete(self, request, pk=None):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên."}, status=status.HTTP_403_FORBIDDEN)

        task_id = pk or request.query_params.get("task_id") or request.data.get("task_id")
        task = get_object_or_404(SupervisionTask, id=task_id, project__supervisor=supervisor)
        task.delete()
        return Response({"message": "Đã xóa nhiệm vụ thành công!"}, status=status.HTTP_200_OK)


# ==============================================================================
# SUPERVISOR BROADCAST ANNOUNCEMENT (Feature 2)
# ==============================================================================

class SupervisorBroadcastAnnouncementAPIView(APIView):
    """
    Giảng viên gửi tin nhắn thông báo chung cho tất cả các nhóm mình hướng dẫn:
    - Nhập tiêu đề và nội dung thông báo chung
    - Tất cả sinh viên trong các nhóm do GV hướng dẫn nhận được thông báo
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response(
                {"message": "Chỉ dành cho giảng viên hướng dẫn."},
                status=status.HTTP_403_FORBIDDEN,
            )

        title = request.data.get("title", "").strip()
        message = request.data.get("message", "").strip()

        if not title:
            return Response(
                {"message": "Tiêu đề thông báo không được để trống."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not message:
            return Response(
                {"message": "Nội dung thông báo không được để trống."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        recipient_users = set()

        # 1. From accepted supervised groups
        supervised_groups = SupervisorOfStudentGroup.objects.filter(
            supervisor=supervisor, status="accepted"
        ).select_related("group")

        for sg in supervised_groups:
            group = sg.group
            if group:
                for member in group.members.select_related("student__user").all():
                    if member.student and member.student.user:
                        recipient_users.add(member.student.user)
                if group.student_1 and group.student_1.user:
                    recipient_users.add(group.student_1.user)
                if group.student_2 and group.student_2.user:
                    recipient_users.add(group.student_2.user)

        # 2. From individual graduation projects
        grad_projects = GraduationProject.objects.filter(
            supervisor=supervisor
        ).select_related("student__user")

        for gp in grad_projects:
            if gp.student and gp.student.user:
                recipient_users.add(gp.student.user)

        if not recipient_users:
            return Response(
                {"message": "Bạn hiện chưa có sinh viên hoặc nhóm nào được phân công hướng dẫn."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        supervisor_name = supervisor.user.get_full_name() or supervisor.user.username
        full_title = f"[Thông báo GVHD {supervisor_name}] {title}"

        for u in recipient_users:
            NotificationService.create_notification(
                user=u,
                notification_type="general",
                title=full_title,
                message=message,
                action_url="/student/dashboard",
                send_email=True,
            )

        return Response(
            {
                "message": f"Đã gửi thông báo chung thành công tới {len(recipient_users)} sinh viên.",
                "recipient_count": len(recipient_users),
            },
            status=status.HTTP_200_OK,
        )



# ==============================================================================
# REVIEWER DASHBOARD & EVALUATION
# ==============================================================================

class ReviewerAssignedProjectsAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên phản biện."}, status=status.HTTP_403_FORBIDDEN)

        projects = GraduationProject.objects.filter(reviewer=supervisor).select_related(
            "student__user",
            "supervisor__user",
            "council",
            "final_grade_summary"
        )
        return Response(GraduationProjectDetailSerializer(projects, many=True).data, status=status.HTTP_200_OK)


class ReviewerSubmitEvaluationAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên phản biện."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.data.get("project_id")
        reviewer_score = request.data.get("reviewer_score")
        reviewer_feedback = request.data.get("reviewer_feedback", "").strip()
        reviewer_verdict = request.data.get("reviewer_verdict", "APPROVED")
        if reviewer_verdict not in ["APPROVED", "CONDITIONAL", "REJECTED"]:
            reviewer_verdict = "APPROVED"

        try:
            score = float(reviewer_score)
            if score < 0.0 or score > 10.0:
                return Response({"reviewer_score": ["Điểm phản biện phải từ 0.0 đến 10.0"]}, status=status.HTTP_400_BAD_REQUEST)
        except (ValueError, TypeError):
            return Response({"reviewer_score": ["Điểm số không hợp lệ."]}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id, reviewer=supervisor)

        with transaction.atomic():
            project.reviewer_score = score
            project.reviewer_feedback = reviewer_feedback
            project.reviewer_verdict = reviewer_verdict
            project.save()

            # Update Final Grade
            summary, _ = FinalGradeSummary.objects.get_or_create(project=project)
            summary.reviewer_score = score
            policy = EvaluationPolicy.objects.filter(batch=project.batch).first()
            summary.calculate_and_save(policy=policy)

        verdict_text = getattr(project, "get_reviewer_verdict_display", lambda: reviewer_verdict)()
        return Response({
            "message": f"Đã lưu nhận xét và kết luận phản biện ({verdict_text})!",
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# COUNCIL LIVE DEFENSE & LIVE GRADING
# ==============================================================================

class CouncilLiveDefenseSessionAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        council_member = CouncilMember.objects.filter(user=user).select_related("council__batch").first()

        if not council_member:
            return Response({"detail": "Bạn không thuộc Hội đồng bảo vệ nào trong kỳ này."}, status=status.HTTP_404_NOT_FOUND)

        council = council_member.council
        all_members = list(council.members.select_related("user", "supervisor").all())

        projects = GraduationProject.objects.filter(council=council).select_related(
            "student__user",
            "supervisor__user",
            "reviewer__user",
            "reviewer__user",
            "final_grade_summary"
        )

        from .utils.vietnamese_sort import sort_by_vietnamese_name
        projects_list = list(projects)
        projects_list = sort_by_vietnamese_name(
            projects_list,
            key_extractor=lambda p: p.student.user.get_full_name() or p.student.user.username
        )

        # Get scores submitted by this member and all scores in council
        member_scores = {getattr(s, "project_id", getattr(s.project, "id", None)): s for s in CouncilLiveScore.objects.filter(member=council_member)}
        all_live_scores = list(CouncilLiveScore.objects.filter(council=council).select_related("member__user"))

        # Group scores by project_id
        scores_by_project = {}
        for s in all_live_scores:
            p_id = s.project_id
            if p_id not in scores_by_project:
                scores_by_project[p_id] = {}
            scores_by_project[p_id][s.member_id] = s

        projects_data = []
        for p in projects_list:
            p_data = dict(GraduationProjectDetailSerializer(p).data)
            my_score = member_scores.get(getattr(p, "id", None))
            p_data["my_score"] = CouncilLiveScoreSerializer(my_score).data if my_score else None

            # Build scoring status breakdown for this project
            p_scores = scores_by_project.get(p.id, {})
            eligible_count = 0
            submitted_count = 0
            members_breakdown = []

            for m in all_members:
                is_sup = bool(m.supervisor_id and m.supervisor_id == p.supervisor_id)
                m_score = p_scores.get(m.id)
                has_sub = m_score is not None
                if not is_sup:
                    eligible_count += 1
                    if has_sub:
                        submitted_count += 1

                members_breakdown.append({
                    "member_id": m.id,
                    "name": m.user.get_full_name() or m.user.username,
                    "role": m.get_role_display(),
                    "role_code": m.role,
                    "is_supervisor": is_sup,
                    "has_submitted": has_sub,
                    "total_score": m_score.total_score if m_score else None,
                })

            p_conflicts = CouncilConflictService.check_project_assignment(council, p)
            p_data["conflicts"] = p_conflicts
            p_data["has_conflict"] = len(p_conflicts) > 0

            p_data["scoring_summary"] = {
                "total_eligible_members": eligible_count,
                "submitted_count": submitted_count,
                "pending_count": max(0, eligible_count - submitted_count),
                "is_fully_graded": eligible_count > 0 and submitted_count >= eligible_count,
                "members_breakdown": members_breakdown
            }
            projects_data.append(p_data)

        role_display = getattr(council_member, "get_role_display", lambda: council_member.role)()
        session_time_display = getattr(council, "get_session_time_display", lambda: council.session_time)()
        can_lock = role_display in ["Chủ tịch hội đồng", "Ủy viên, Thư ký"] or user.is_staff
        council_conflict_info = CouncilConflictService.check_council_conflicts(council_id=council.id)

        return Response({
            "council": {
                "id": getattr(council, "id", None),
                "council_number": council.council_number,
                "council_name": council.council_name,
                "session_date": council.session_date,
                "session_time": session_time_display,
                "defense_room": council.defense_room,
                "my_role": role_display,
                "is_locked": getattr(council, "is_locked", False),
                "locked_at": council.locked_at if getattr(council, "is_locked", False) else None,
                "locked_by": council.locked_by.get_full_name() or council.locked_by.username if getattr(council, "is_locked", False) and council.locked_by else None,
                "can_lock": can_lock,
                "my_role_code": council_member.role,
                "current_defending_project_id": council.current_defending_project_id,
                "has_conflict": council_conflict_info.get("has_conflict", False),
                "total_conflicts": council_conflict_info.get("total_conflicts", 0),
                "conflicts": council_conflict_info.get("conflicts", []),
            },
            "members": [
                {
                    "id": m.id,
                    "name": m.user.get_full_name() or m.user.username,
                    "role": m.get_role_display(),
                    "role_code": m.role,
                }
                for m in all_members
            ],
            "projects": projects_data
        }, status=status.HTTP_200_OK)


class CouncilSubmitScoreAPIView(APIView):
    permission_classes = [IsAuthenticated, IsSupervisorOrCommitteeMember]

    @retry_on_db_lock(max_retries=5, initial_delay=0.05, backoff_factor=1.5)
    def post(self, request):
        user = request.user
        if getattr(user, "user_type", "") == "student":
            return Response({"detail": "Sinh viên không có quyền truy cập hoặc tự nhập điểm Hội đồng bảo vệ."}, status=status.HTTP_403_FORBIDDEN)

        council_member = CouncilMember.objects.filter(user=user).first()
        if not council_member:
            return Response({"detail": "Bạn không phải thành viên Hội đồng bảo vệ."}, status=status.HTTP_403_FORBIDDEN)

        if council_member.council and getattr(council_member.council, "is_locked", False):
            return Response(
                {"detail": "Hội đồng đã khóa điểm. Toàn bộ điểm số ở trạng thái chỉ đọc (Read-only), không thể chỉnh sửa."},
                status=status.HTTP_403_FORBIDDEN
            )

        project_id = request.data.get("project_id")
        p_score = float(request.data.get("score_presentation", 0.0))
        c_score = float(request.data.get("score_content", 0.0))
        q_score = float(request.data.get("score_qa", 0.0))
        d_score = float(request.data.get("score_demo", 0.0))
        comments = request.data.get("comments", "").strip()

        # Strict checks: Supervisor cannot grade their own student in council
        project = get_object_or_404(GraduationProject, id=project_id, council=council_member.council)
        if council_member.supervisor and project.supervisor == council_member.supervisor:
            return Response({"detail": "Vi phạm quy chế: Giảng viên hướng dẫn không được chấm điểm Hội đồng cho sinh viên của mình."}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            project = GraduationProject.objects.select_for_update().get(id=project_id, council=council_member.council)
            live_score, _ = CouncilLiveScore.objects.update_or_create(
                project=project,
                member=council_member,
                defaults={
                    "council": council_member.council,
                    "score_presentation": p_score,
                    "score_content": c_score,
                    "score_qa": q_score,
                    "score_demo": d_score,
                    "comments": comments
                }
            )

            # Re-calculate average council score
            all_scores = list(CouncilLiveScore.objects.filter(project=project))
            if all_scores:
                avg_council = round(sum(s.total_score for s in all_scores) / len(all_scores), 2)
            else:
                avg_council = 0.0

            # Update FinalGradeSummary
            summary, _ = FinalGradeSummary.objects.select_for_update().get_or_create(project=project)
            summary.supervisor_score = project.supervisor_score
            summary.reviewer_score = project.reviewer_score
            summary.council_avg_score = avg_council

            policy = EvaluationPolicy.objects.filter(batch=project.batch).first()
            summary.calculate_and_save(policy=policy)

            # If all evaluations exist, mark PASSED/FAILED
            if summary.final_score_10 is not None:
                project.status = "PASSED" if summary.is_passed else "FAILED"
                project.save()

            # Automatic notification when External council member submits evaluation score
            if council_member.role == "EXTERNAL_MEMBER":
                ext_name = council_member.user.get_full_name() or council_member.user.username
                # Notify Chair & Secretary of council
                leaders = CouncilMember.objects.filter(
                    council=council_member.council,
                    role__in=["CHAIR", "SECRETARY"]
                ).select_related("user")
                for leader in leaders:
                    if leader.user != user:
                        NotificationService.create_notification(
                            user=leader.user,
                            notification_type="evaluation_completed",
                            title="Chuyên gia ngoài đã nộp phiếu đánh giá",
                            message=f"Chuyên gia ngoài (Ủy viên ngoài trường) {ext_name} đã hoàn tất chấm điểm cho sinh viên {project.student.user.get_full_name()} ({project.student.registration_no}): {live_score.total_score}đ.",
                        )
                # Notify Student
                NotificationService.create_notification(
                    user=project.student.user,
                    notification_type="evaluation_completed",
                    title="Chuyên gia ngoài đã nộp phiếu đánh giá",
                    message=f"Chuyên gia ngoài (Ủy viên ngoài trường) {ext_name} đã nộp phiếu đánh giá bảo vệ cho đồ án của bạn: {live_score.total_score}đ.",
                )

        return Response({
            "message": f"Đã chấm điểm thành công cho SV {project.student.registration_no}: {live_score.total_score}đ (Điểm TB HĐ: {summary.council_avg_score}đ - Tổng kết: {summary.final_score_10}đ)",
            "live_score": CouncilLiveScoreSerializer(live_score).data,
            "final_grade": FinalGradeSummarySerializer(summary).data
        }, status=status.HTTP_200_OK)


class CouncilToggleLockAPIView(APIView):
    """Khóa hoặc mở khóa bảng điểm hội đồng"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        council_id = request.data.get("council_id")
        lock = request.data.get("lock", True)

        council = get_object_or_404(DefenseCouncil, id=council_id)
        membership = CouncilMember.objects.filter(council=council, user=user).first()
        is_chair_or_secretary = membership and membership.role in ["CHAIR", "SECRETARY"]
        is_admin = user.is_staff or getattr(user, "user_type", None) == "committee_member"

        if not is_chair_or_secretary and not is_admin:
            return Response(
                {"detail": "Chỉ Chủ tịch hoặc Thư ký hội đồng mới có quyền khóa hoặc mở khóa điểm."},
                status=status.HTTP_403_FORBIDDEN
            )

        council.is_locked = bool(lock)
        council.locked_at = timezone.now() if lock else None
        council.locked_by = user if lock else None
        council.save(update_fields=["is_locked", "locked_at", "locked_by"])

        action = "khóa" if lock else "mở khóa"
        return Response({
            "message": f"Đã {action} điểm hội đồng {council.council_name} thành công!",
            "is_locked": council.is_locked,
            "locked_at": council.locked_at,
            "locked_by": user.get_full_name() or user.username if lock else None
        }, status=status.HTTP_200_OK)


class CouncilChairSetDefenseStatusAPIView(APIView):
    """Chủ tịch hội đồng điều hành buổi bảo vệ: Chuyển trạng thái đồ án sang 'Đang bảo vệ' (In Progress / DEFENDING)"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        council_member = CouncilMember.objects.filter(user=user).select_related("council").first()
        if not council_member:
            return Response({"detail": "Bạn không thuộc Hội đồng bảo vệ nào."}, status=status.HTTP_403_FORBIDDEN)

        if council_member.role not in ["CHAIR", "SECRETARY"]:
            return Response({"detail": "Chỉ Chủ tịch hội đồng (hoặc Thư ký) mới có quyền điều hành trạng thái buổi bảo vệ."}, status=status.HTTP_403_FORBIDDEN)

        council = council_member.council
        project_id = request.data.get("project_id")
        defense_status = request.data.get("defense_status", "DEFENDING")  # DEFENDING, DEFENDED, WAITING

        if not project_id:
            return Response({"detail": "project_id là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id, council=council)

        with transaction.atomic():
            if defense_status == "DEFENDING":
                # If there's another project currently defending in council, mark it DEFENDED
                GraduationProject.objects.filter(council=council, defense_status="DEFENDING").exclude(id=project.id).update(defense_status="DEFENDED")

                project.defense_status = "DEFENDING"
                project.status = "DEFENDING"
                project.save()

                council.current_defending_project = project
                council.save(update_fields=["current_defending_project"])

                status_msg = f"Đã chuyển sinh viên {project.student.user.get_full_name()} ({project.student.registration_no}) sang trạng thái 'Đang bảo vệ'."
            elif defense_status == "DEFENDED":
                project.defense_status = "DEFENDED"
                project.save(update_fields=["defense_status"])

                if council.current_defending_project_id == project.id:
                    council.current_defending_project = None
                    council.save(update_fields=["current_defending_project"])

                status_msg = f"Đã hoàn tất phần bảo vệ cho sinh viên {project.student.user.get_full_name()}."
            else:  # WAITING
                project.defense_status = "WAITING"
                if project.status == "DEFENDING":
                    project.status = "DEFENSE_READY"
                project.save()

                if council.current_defending_project_id == project.id:
                    council.current_defending_project = None
                    council.save(update_fields=["current_defending_project"])

                status_msg = f"Đã đặt lại trạng thái chờ bảo vệ cho sinh viên {project.student.user.get_full_name()}."

            AuditLog.objects.create(
                user=user,
                action_type="evaluation_update",
                description=f"Hội đồng {council.council_name}: {status_msg}"
            )

        return Response({
            "message": status_msg,
            "current_defending_project_id": council.current_defending_project_id,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class CouncilSecretaryRemindScoringAPIView(APIView):
    """Thư ký (hoặc Chủ tịch) nhắc nhở các thành viên hội đồng chưa nộp điểm bảo vệ"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        council_member = CouncilMember.objects.filter(user=user).select_related("council").first()
        if not council_member:
            return Response({"detail": "Bạn không phải thành viên Hội đồng bảo vệ."}, status=status.HTTP_403_FORBIDDEN)

        if council_member.role not in ["SECRETARY", "CHAIR"]:
            return Response({"detail": "Chỉ Thư ký hoặc Chủ tịch hội đồng mới có quyền gửi nhắc nhở nộp điểm."}, status=status.HTTP_403_FORBIDDEN)

        council = council_member.council
        project_id = request.data.get("project_id")
        if not project_id:
            if council.current_defending_project_id:
                project_id = council.current_defending_project_id
            else:
                return Response({"detail": "Vui lòng chọn đề tài/sinh viên cần nhắc nộp điểm."}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(GraduationProject, id=project_id, council=council)

        # Find eligible members (exclude supervisor of this student)
        eligible_members = [
            m for m in council.members.select_related("user").all()
            if not (m.supervisor_id and m.supervisor_id == project.supervisor_id)
        ]

        # Find members who already submitted score
        submitted_member_ids = set(CouncilLiveScore.objects.filter(project=project).values_list("member_id", flat=True))

        pending_members = [m for m in eligible_members if m.id not in submitted_member_ids]

        reminded_list = []
        for m in pending_members:
            title = f"🔔 [Hội đồng #{council.council_number}] Nhắc nhở nộp điểm bảo vệ"
            msg = (
                f"Thư ký Hội đồng {council.council_name} kính nhắc Thầy/Cô {m.user.get_full_name()} "
                f"chưa nộp điểm bảo vệ cho SV {project.student.user.get_full_name()} (MSSV: {project.student.registration_no}) "
                f"- Đề tài: '{project.topic_title_vi}'. Kính đề nghị Thầy/Cô hoàn tất chấm điểm."
            )
            try:
                NotificationService.create_notification(
                    user=m.user,
                    notification_type="general",
                    title=title,
                    message=msg,
                    action_url="/utc-live-defense",
                    send_email=True
                )
            except Exception as e:
                logger.warning("Could not send reminder notification: %s", e)

            reminded_list.append({
                "id": m.id,
                "name": m.user.get_full_name() or m.user.username,
                "role": m.get_role_display(),
                "email": m.user.email
            })

        names_str = ", ".join(m["name"] for m in reminded_list) if reminded_list else "Không có thành viên nào chưa nộp"
        success_message = (
            f"Đã gửi thông báo nhắc nhở nộp điểm tới {len(reminded_list)} thành viên: {names_str}."
            if reminded_list
            else "Tất cả các thành viên hội đồng đã nộp đủ điểm cho sinh viên này!"
        )

        return Response({
            "success": True,
            "message": success_message,
            "reminded_members": reminded_list,
            "already_submitted_count": len(submitted_member_ids),
            "pending_count": len(pending_members)
        }, status=status.HTTP_200_OK)


class CouncilScheduleDefenseAPIView(APIView):
    """
    API for Committee / Admin to schedule defense session for a council.
    Updates session_date, session_time, defense_room, and dispatches UTC HTML email
    notifications to students, supervisors, and council members.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request, council_id):
        user = request.user
        if not (user.is_staff or getattr(user, "user_type", None) in ["admin", "committee"]):
            return Response({"detail": "Chỉ Ban chủ nhiệm hoặc Admin mới có quyền xếp lịch bảo vệ."}, status=status.HTTP_403_FORBIDDEN)

        council = get_object_or_404(DefenseCouncil, id=council_id)
        session_date = request.data.get("session_date")
        session_time = request.data.get("session_time")
        defense_room = request.data.get("defense_room", "").strip()

        if session_date:
            council.session_date = session_date
        if session_time:
            council.session_time = session_time
        if defense_room:
            council.defense_room = defense_room

        council.save()

        # Send UTC branded email notifications to all parties (students, supervisors, members)
        NotificationService.notify_defense_scheduled_emails(council)

        return Response({
            "message": f"Xếp lịch bảo vệ thành công cho Hội đồng số {council.council_number} và đã kích hoạt email thông báo nhận diện UTC!",
            "council": {
                "id": council.id,
                "council_number": council.council_number,
                "session_date": str(council.session_date),
                "session_time": council.session_time,
                "defense_room": council.defense_room,
            }
        }, status=status.HTTP_200_OK)


# ==============================================================================
# COUNCIL CONFLICT OF INTEREST & ASSIGNMENT MANAGEMENT
# ==============================================================================

class CouncilConflictCheckAPIView(APIView):
    """
    API kiểm tra xung đột lợi ích (Conflict of Interest) trong phân công hội đồng bảo vệ.
    Hỗ trợ kiểm tra toàn bộ hội đồng trong đợt hoặc một hội đồng cụ thể.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        council_id = request.query_params.get("council_id")
        batch_id = request.query_params.get("batch_id")

        if council_id:
            try:
                council_id = int(council_id)
            except ValueError:
                return Response({"detail": "council_id phải là số nguyên."}, status=status.HTTP_400_BAD_REQUEST)

        if batch_id:
            try:
                batch_id = int(batch_id)
            except ValueError:
                return Response({"detail": "batch_id phải là số nguyên."}, status=status.HTTP_400_BAD_REQUEST)

        data = CouncilConflictService.check_council_conflicts(council_id=council_id, batch_id=batch_id)
        return Response(data, status=status.HTTP_200_OK)


class CouncilAssignProjectAPIView(APIView):
    """
    API phân công đề tài vào hội đồng bảo vệ có kiểm tra cảnh báo xung đột lợi ích.
    """
    permission_classes = [IsAuthenticated]

    def post(self, request):
        project_id = request.data.get("project_id")
        council_id = request.data.get("council_id")
        force = request.data.get("force", False)
        if isinstance(force, str):
            force = force.lower() in ("true", "1")
        else:
            force = bool(force)

        if not project_id:
            return Response({"detail": "project_id là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        project = get_object_or_404(
            GraduationProject.objects.select_related("student__user", "supervisor__user", "reviewer__user"),
            id=project_id
        )

        # Unassign if council_id is None or 0
        if council_id is None or council_id == "" or council_id == 0:
            old_council_name = project.council.council_name if project.council else "Không"
            project.council = None
            project.save(update_fields=["council"])

            AuditLog.objects.create(
                user=request.user,
                action_type="status_change",
                description=f"Hủy phân công đề tài {project.student.registration_no} khỏi hội đồng {old_council_name}."
            )
            return Response({
                "success": True,
                "message": f"Đã hủy phân công đề tài khỏi hội đồng.",
                "project_id": project.id,
                "council_id": None
            }, status=status.HTTP_200_OK)

        council = get_object_or_404(DefenseCouncil, id=council_id)

        # Boundary check: Quy chế UTC tối đa 12 sinh viên trong một buổi bảo vệ (max_limit=12)
        current_assigned_count = council.projects.exclude(id=project.id).count()
        if current_assigned_count >= 12:
            return Response({
                "success": False,
                "error": "limit_exceeded",
                "max_limit": 12,
                "current_count": current_assigned_count,
                "detail": f"Hội đồng '{council.council_name}' đã đạt giới hạn tối đa 12 sinh viên trong một buổi bảo vệ. Không thể xếp thêm sinh viên thứ {current_assigned_count + 1}.",
                "message": f"Hội đồng '{council.council_name}' đã đạt giới hạn tối đa 12 sinh viên trong một buổi bảo vệ. Không thể xếp thêm sinh viên thứ {current_assigned_count + 1}."
            }, status=status.HTTP_400_BAD_REQUEST)

        conflicts = CouncilConflictService.check_project_assignment(council, project)

        if conflicts and not force:
            return Response({
                "success": False,
                "has_conflict": True,
                "conflicts_count": len(conflicts),
                "conflicts": conflicts,
                "message": conflicts[0].get("message", f"Không thể xếp sinh viên vào Hội đồng có GVHD tham gia chấm. ({len(conflicts)} vi phạm)")
            }, status=status.HTTP_400_BAD_REQUEST)

        try:
            project.council = council
            project.save(update_fields=["council"])
        except Exception as e:
            return Response({
                "success": False,
                "error": "constraint_violation",
                "detail": f"Lỗi ràng buộc CSDL: {str(e)}",
                "message": f"Không thể xếp sinh viên vào hội đồng do vi phạm ràng buộc CSDL: {str(e)}"
            }, status=status.HTTP_400_BAD_REQUEST)

        AuditLog.objects.create(
            user=request.user,
            action_type="status_change",
            description=f"Phân công đề tài {project.student.registration_no} vào {council.council_name}."
        )

        return Response({
            "success": True,
            "message": f"Đã phân công đề tài của sinh viên {project.student.user.get_full_name()} vào {council.council_name}.",
            "project_id": project.id,
            "council_id": council.id,
            "council_name": council.council_name,
            "has_conflict": len(conflicts) > 0,
            "conflicts": conflicts
        }, status=status.HTTP_200_OK)


class CouncilAssignMemberAPIView(APIView):
    """
    API thêm hoặc cập nhật thành viên vào hội đồng bảo vệ có kiểm tra cảnh báo xung đột lợi ích
    và tối ưu hóa phân bổ hướng nghiên cứu (TC_039: Cân bằng Hội đồng: Ưu tiên gán phản biện có cùng lĩnh vực với Đề tài).
    Quy chế: GV hướng dẫn KHÔNG nằm trong hội đồng/phản biện của chính sinh viên đó.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        """Gợi ý danh sách giảng viên phản biện / ủy viên được tối ưu theo hướng nghiên cứu của các đề tài trong hội đồng"""
        council_id = request.query_params.get("council_id")
        if not council_id:
            return Response({"detail": "council_id là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        council = get_object_or_404(DefenseCouncil, id=council_id)
        projects = list(council.projects.all().select_related("topic_category", "supervisor__user"))
        supervisors = list(Supervisor.objects.all().select_related("user"))

        existing_user_ids = set(council.members.values_list("user_id", flat=True))

        recommendations = []
        for s in supervisors:
            # Check COI
            conflicts = CouncilConflictService.check_member_assignment(council, s.user, s)
            has_coi = len(conflicts) > 0

            # Calculate match score across all projects in council
            scores = [CouncilStructureService.calculate_topic_match_score(s, p) for p in projects]
            avg_score = round(sum(scores) / len(scores), 1) if scores else 0.0
            max_score = max(scores, default=0.0)

            recommendations.append({
                "supervisor_id": s.id,
                "user_id": s.user.id,
                "full_name": s.user.get_full_name() or s.user.username,
                "academic_title": s.academic_title or "ThS",
                "department": s.department_name,
                "research_interest": s.research_interest,
                "topic_match_score": max_score,
                "avg_match_score": avg_score,
                "is_topic_matched": max_score >= 3.0,
                "has_coi": has_coi,
                "is_already_member": s.user.id in existing_user_ids,
                "conflicts": conflicts
            })

        recommendations.sort(key=lambda x: (not x["has_coi"], not x["is_already_member"], x["topic_match_score"]), reverse=True)

        return Response({
            "council_id": council.id,
            "council_name": council.council_name,
            "projects_count": len(projects),
            "recommendations": recommendations[:20]
        }, status=status.HTTP_200_OK)

    def post(self, request):
        council_id = request.data.get("council_id")
        user_id = request.data.get("user_id")
        supervisor_id = request.data.get("supervisor_id")
        role = request.data.get("role", "MEMBER")
        auto_assign = request.data.get("auto_assign", False)
        optimize_topics = request.data.get("optimize_topics", False)

        if isinstance(auto_assign, str):
            auto_assign = auto_assign.lower() in ("true", "1", "yes")
        if isinstance(optimize_topics, str):
            optimize_topics = optimize_topics.lower() in ("true", "1", "yes")

        if not council_id:
            return Response({"detail": "council_id là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        council = get_object_or_404(DefenseCouncil, id=council_id)
        council_projects = list(council.projects.all().select_related("topic_category", "supervisor__user"))

        # Thuật toán tự động tối ưu hóa gán phản biện theo hướng nghiên cứu (TC_039)
        if (auto_assign or optimize_topics or not (user_id or supervisor_id)) and council_projects:
            all_supervisors = list(Supervisor.objects.select_related("user").all())
            existing_sup_ids = set(council.members.filter(supervisor__isnull=False).values_list("supervisor_id", flat=True))

            best_sup = None
            best_score = -1.0

            for cand in all_supervisors:
                if cand.id in existing_sup_ids:
                    continue
                coi = CouncilConflictService.check_member_assignment(council, cand.user, cand)
                if coi:
                    continue
                cand_scores = [CouncilStructureService.calculate_topic_match_score(cand, p) for p in council_projects]
                max_cand_score = max(cand_scores, default=0.0)
                if max_cand_score > best_score:
                    best_score = max_cand_score
                    best_sup = cand

            if best_sup:
                user = best_sup.user
                supervisor = best_sup
                role = role or "REVIEWER"
            else:
                return Response({
                    "success": False,
                    "detail": "Không tìm thấy giảng viên phù hợp không có xung đột lợi ích để phân công."
                }, status=status.HTTP_400_BAD_REQUEST)
        else:
            if not user_id and not supervisor_id:
                return Response({"detail": "council_id và user_id (hoặc supervisor_id) là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

            user = None
            supervisor = None
            if supervisor_id:
                supervisor = Supervisor.objects.filter(id=supervisor_id).first()
                if supervisor:
                    user = supervisor.user

            if not user and user_id:
                user = CustomUser.objects.filter(id=user_id).first()
                if user:
                    supervisor = Supervisor.objects.filter(user=user).first()
                if not supervisor:
                    supervisor = Supervisor.objects.filter(id=user_id).first()
                    if supervisor and not user:
                        user = supervisor.user

            if not user:
                return Response({"detail": "Không tìm thấy thông tin giảng viên/người dùng."}, status=status.HTTP_404_NOT_FOUND)

        # Cross-check conflict of interest: GVHD cannot evaluate/review their own students in council
        conflicts = CouncilConflictService.check_member_assignment(council, user, supervisor)

        if conflicts:
            return Response({
                "success": False,
                "has_conflict": True,
                "conflicts_count": len(conflicts),
                "conflicts": conflicts,
                "message": conflicts[0].get("message", f"Xung đột lợi ích: Giảng viên hướng dẫn không được làm thành viên/phản biện trong hội đồng của chính sinh viên đó.")
            }, status=status.HTTP_400_BAD_REQUEST)

        # Calculate topic match score (TC_039: Cân bằng hội đồng & Ưu tiên cùng lĩnh vực)
        topic_match_score = 0.0
        if supervisor and council_projects:
            scores = [CouncilStructureService.calculate_topic_match_score(supervisor, p) for p in council_projects]
            topic_match_score = max(scores, default=0.0)

            # Gán làm GVPB (Reviewer) cho các đề tài trong hội đồng nếu vai trò là REVIEWER hoặc MEMBER
            if role in ["REVIEWER", "MEMBER"]:
                for p in council_projects:
                    if not p.reviewer or CouncilStructureService.calculate_topic_match_score(supervisor, p) >= 3.0:
                        p.reviewer = supervisor
                        p.save(update_fields=["reviewer"])

        member, created = CouncilMember.objects.update_or_create(
            council=council,
            user=user,
            defaults={
                "role": role,
                "supervisor": supervisor
            }
        )

        AuditLog.objects.create(
            user=request.user,
            action_type="status_change",
            description=f"{'Thêm' if created else 'Cập nhật'} ủy viên {user.get_full_name()} ({member.get_role_display()}) vào {council.council_name}."
        )

        return Response({
            "success": True,
            "message": f"Đã phân công Thầy/Cô {user.get_full_name()} vào {council.council_name} ({member.get_role_display()}).",
            "member_id": member.id,
            "council_id": council.id,
            "role": member.role,
            "role_display": member.get_role_display(),
            "topic_match_score": topic_match_score,
            "is_topic_matched": topic_match_score >= 3.0,
            "has_conflict": len(conflicts) > 0,
            "conflicts": conflicts
        }, status=status.HTTP_200_OK)


# Alias AssignMemberAPIView for TC_039
AssignMemberAPIView = CouncilAssignMemberAPIView


# ==============================================================================
# GLOBAL SEARCH API
# ==============================================================================

class GlobalSearchAPIView(APIView):
    """
    Thanh tìm kiếm toàn cục (Global Search) ở Top Header.
    Tìm kiếm tức thời đa đối tượng: Đề tài, Sinh viên, Giảng viên, Hội đồng, Biểu mẫu.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from django.db.models import Q

        q = request.query_params.get("q", "").strip()
        search_type = request.query_params.get("type", "all").lower()

        if not q or len(q) < 2:
            return Response({
                "query": q,
                "total_results": 0,
                "results": [],
                "categories": {
                    "projects": 0,
                    "faculty": 0,
                    "students": 0,
                    "councils": 0,
                    "documents": 0
                }
            }, status=status.HTTP_200_OK)

        results = []

        # 1. Projects (GraduationProject & Project)
        if search_type in ["all", "projects"]:
            # Graduation Projects (UTC)
            grad_projects = GraduationProject.objects.filter(
                Q(topic_title_vi__icontains=q) |
                Q(topic_title_en__icontains=q) |
                Q(student__registration_no__icontains=q) |
                Q(student__user__first_name__icontains=q) |
                Q(student__user__last_name__icontains=q) |
                Q(supervisor__user__first_name__icontains=q) |
                Q(supervisor__user__last_name__icontains=q)
            ).select_related("student__user", "supervisor__user", "council")[:8]

            for p in grad_projects:
                s_name = p.student.user.get_full_name() or p.student.user.username
                results.append({
                    "id": f"grad_project_{p.id}",
                    "item_id": p.id,
                    "type": "project",
                    "type_label": "Đồ án Tốt nghiệp",
                    "icon": "🎓",
                    "title": p.topic_title_vi or p.topic_title_en,
                    "subtitle": f"SV: {s_name} ({p.student.registration_no}) • Lớp: {p.student.course_class or 'K62'}",
                    "extra_info": f"HĐ: {p.council.council_name if p.council else 'Chưa gán'} • Trạng thái: {p.get_status_display()}",
                    "action_url": f"/student/dashboard?tab=overview",
                })

            # Standard FYP Projects
            fyp_projects = Project.objects.filter(
                Q(project_name__icontains=q) |
                Q(project_description__icontains=q) |
                Q(language__icontains=q)
            ).select_related("project_category")[:5]

            for p in fyp_projects:
                results.append({
                    "id": f"fyp_project_{p.id}",
                    "item_id": p.id,
                    "type": "project",
                    "type_label": "Đề tài FYP",
                    "icon": "📁",
                    "title": p.project_name,
                    "subtitle": f"Ngôn ngữ: {p.language or 'N/A'} • Danh mục: {p.project_category.category_name if p.project_category else 'N/A'}",
                    "extra_info": p.project_description[:80] + "..." if p.project_description and len(p.project_description) > 80 else p.project_description,
                    "action_url": f"/student/dashboard?tab=project",
                })

        # 2. Faculty / Supervisors
        if search_type in ["all", "faculty"]:
            supervisors = Supervisor.objects.filter(
                Q(user__first_name__icontains=q) |
                Q(user__last_name__icontains=q) |
                Q(user__username__icontains=q) |
                Q(user__email__icontains=q) |
                Q(supervisor_id__icontains=q) |
                Q(department_name__icontains=q)
            ).select_related("user")[:6]

            for s in supervisors:
                name = s.user.get_full_name() or s.user.username
                results.append({
                    "id": f"supervisor_{s.id}",
                    "item_id": s.id,
                    "type": "faculty",
                    "type_label": "Giảng viên",
                    "icon": "👨‍🏫",
                    "title": f"{s.academic_title or 'ThS/TS'}. {name}",
                    "subtitle": f"Bộ môn: {s.department_name or 'CNTT'} • Email: {s.user.email or 'N/A'}",
                    "extra_info": f"Mã GV: {s.supervisor_id}",
                    "action_url": f"/supervisor/dashboard",
                })

        # 3. Students
        if search_type in ["all", "students"]:
            students = Student.objects.filter(
                Q(user__first_name__icontains=q) |
                Q(user__last_name__icontains=q) |
                Q(user__username__icontains=q) |
                Q(user__email__icontains=q) |
                Q(registration_no__icontains=q) |
                Q(department__icontains=q)
            ).select_related("user")[:6]

            for st in students:
                name = st.user.get_full_name() or st.user.username
                results.append({
                    "id": f"student_{st.id}",
                    "item_id": st.id,
                    "type": "student",
                    "type_label": "Sinh viên",
                    "icon": "👨‍🎓",
                    "title": name,
                    "subtitle": f"MSSV: {st.registration_no} • Khoa: {st.department or 'CNTT'}",
                    "extra_info": f"Email: {st.user.email or 'N/A'}",
                    "action_url": f"/student/dashboard",
                })

        # 4. Councils & Panels
        if search_type in ["all", "councils"]:
            councils = DefenseCouncil.objects.filter(
                Q(council_name__icontains=q) |
                Q(defense_room__icontains=q)
            ).select_related("batch")[:6]

            for c in councils:
                results.append({
                    "id": f"council_{c.id}",
                    "item_id": c.id,
                    "type": "council",
                    "type_label": "Hội đồng Bảo vệ",
                    "icon": "🏛️",
                    "title": f"{c.council_name} (HĐ #{c.council_number})",
                    "subtitle": f"Phòng: {c.defense_room or 'TBA'} • Ngày: {c.session_date or 'TBA'} ({c.session_time})",
                    "extra_info": f"Đợt: {c.batch.batch_code if c.batch else 'Kỳ hiện tại'}",
                    "action_url": f"/committee_member/dashboard?tab=utc_live_defense",
                })

        # 5. Documents & Templates
        if search_type in ["all", "documents"]:
            doc_reqs = DocumentRequirement.objects.filter(
                Q(title__icontains=q) |
                Q(document_type__icontains=q)
            )[:5]

            for d in doc_reqs:
                results.append({
                    "id": f"doc_{d.id}",
                    "item_id": d.id,
                    "type": "document",
                    "type_label": "Biểu mẫu / Tài liệu",
                    "icon": "📄",
                    "title": d.title,
                    "subtitle": f"Loại: {d.get_document_type_display()} • Hạn nộp: {d.deadline.strftime('%d/%m/%Y %H:%M') if d.deadline else 'Không có'}",
                    "extra_info": f"Học kỳ: {d.semester}",
                    "action_url": f"/student/dashboard?tab=documents",
                })

        # Tally counts by category
        categories_count = {
            "projects": sum(1 for r in results if r["type"] == "project"),
            "faculty": sum(1 for r in results if r["type"] == "faculty"),
            "students": sum(1 for r in results if r["type"] == "student"),
            "councils": sum(1 for r in results if r["type"] == "council"),
            "documents": sum(1 for r in results if r["type"] == "document"),
        }

        return Response({
            "query": q,
            "total_results": len(results),
            "categories": categories_count,
            "results": results
        }, status=status.HTTP_200_OK)


# ==============================================================================
# PHASE 2: ALLOCATION OPTIMIZATION & OVERRIDE APIS
# ==============================================================================

class RunAllocationAlgorithmAPIView(APIView):
    """Khoa / Admin chạy thuật toán tối ưu phân công nguyện vọng đồ án"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        batch_id = request.data.get("batch_id")
        if batch_id:
            batch = get_object_or_404(AcademicBatch, id=batch_id)
        else:
            batch = AcademicBatch.objects.filter(is_active=True).first()
            if not batch:
                return Response({"detail": "Không có đợt đồ án nào đang hoạt động."}, status=status.HTTP_400_BAD_REQUEST)

        stats = ThesisAllocationService.run_allocation(batch)
        allocations = ProposedAllocation.objects.filter(batch=batch).select_related(
            "student__user", "supervisor__user"
        )
        data = ProposedAllocationSerializer(allocations, many=True).data

        return Response({
            "success": True,
            "message": f"Chạy thuật toán tối ưu phân công thành công cho đợt {batch.batch_code}!",
            "stats": stats,
            "allocations": data,
        }, status=status.HTTP_200_OK)


class ProposedAllocationListAPIView(APIView):
    """Xem danh sách đề xuất phân công để Khoa review"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        batch_id = request.query_params.get("batch_id")
        if batch_id:
            batch = get_object_or_404(AcademicBatch, id=batch_id)
        else:
            batch = AcademicBatch.objects.filter(is_active=True).first()
            if not batch:
                return Response({"detail": "Không có đợt đồ án nào đang hoạt động."}, status=status.HTTP_400_BAD_REQUEST)

        allocations = ProposedAllocation.objects.filter(batch=batch).select_related(
            "student__user", "supervisor__user"
        )
        return Response(ProposedAllocationSerializer(allocations, many=True).data, status=status.HTTP_200_OK)


class OverrideAllocationAPIView(APIView):
    """Khoa review / override: Điều chỉnh phân công thủ công"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        allocation_id = request.data.get("allocation_id")
        supervisor_id = request.data.get("supervisor_id")
        reason = request.data.get("reason", "Khoa điều chỉnh thủ công").strip()

        if not allocation_id or not supervisor_id:
            return Response({"detail": "allocation_id và supervisor_id là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        alloc = get_object_or_404(ProposedAllocation, id=allocation_id)
        supervisor = get_object_or_404(Supervisor, id=supervisor_id)

        # Kiểm tra học vị Kỹ sư nếu SV thuộc CT Kỹ sư
        if alloc.student.degree_program == "ENGINEER":
            if not DegreeEligibilityService.is_doctoral_degree(supervisor.academic_title or ""):
                return Response({
                    "detail": f"Không thể phân công: SV {alloc.student.registration_no} thuộc CT Kỹ sư, GVHD bắt buộc có học vị Tiến sĩ trở lên."
                }, status=status.HTTP_400_BAD_REQUEST)

        alloc.supervisor = supervisor
        alloc.is_overridden = True
        alloc.override_reason = reason
        alloc.save(update_fields=["supervisor", "is_overridden", "override_reason"])

        return Response({
            "message": f"Khoa đã điều chỉnh GVHD thành công cho SV {alloc.student.registration_no}!",
            "allocation": ProposedAllocationSerializer(alloc).data
        }, status=status.HTTP_200_OK)


class FinalizeAllocationAPIView(APIView):
    """Khoa chốt phân công: Chuyển toàn bộ đề xuất thành GraduationProject và kích hoạt email"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        batch_id = request.data.get("batch_id")
        if batch_id:
            batch = get_object_or_404(AcademicBatch, id=batch_id)
        else:
            batch = AcademicBatch.objects.filter(is_active=True).first()
            if not batch:
                return Response({"detail": "Không có đợt đồ án nào đang hoạt động."}, status=status.HTTP_400_BAD_REQUEST)

        result = ThesisAllocationService.finalize_allocation(batch, user=request.user)
        return Response({"success": True, **result}, status=status.HTTP_200_OK)


# ==============================================================================
# PHASE 3: TOPIC FORMULATION (DRAFT -> CONFIRM -> APPROVE -> PDF)
# ==============================================================================

class TopicDraftAPIView(APIView):
    """GV và SV cùng xác định đề tài (trạng thái Draft)"""
    permission_classes = [IsAuthenticated]

    def patch(self, request):
        return self._update_topic_draft(request)

    def post(self, request):
        return self._update_topic_draft(request)

    def _update_topic_draft(self, request):
        user = request.user
        project_id = request.data.get("project_id")
        topic_title_vi = request.data.get("topic_title_vi", "").strip()
        topic_title_en = request.data.get("topic_title_en", "").strip()

        if not topic_title_vi:
            return Response({"detail": "Tên đề tài tiếng Việt là bắt buộc."}, status=status.HTTP_400_BAD_REQUEST)

        if user.user_type == "student":
            student = getattr(user, "student_profile", None)
            project = get_object_or_404(GraduationProject, student=student)
        elif user.user_type == "supervisor":
            supervisor = getattr(user, "supervisor_profile", None)
            project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)
        else:
            project = get_object_or_404(GraduationProject, id=project_id)

        project.topic_title_vi = topic_title_vi
        if topic_title_en:
            project.topic_title_en = topic_title_en
        project.status = "TOPIC_DRAFT"
        project.save(update_fields=["topic_title_vi", "topic_title_en", "status"])

        return Response({
            "message": "Đã cập nhật tên đề tài (Trạng thái Draft).",
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class SupervisorConfirmTopicAPIView(APIView):
    """GV xác nhận đề tài (sau khi thống nhất với SV) để trình Khoa/Ban duyệt"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên hướng dẫn."}, status=status.HTTP_403_FORBIDDEN)

        project_id = request.data.get("project_id")
        project = get_object_or_404(GraduationProject, id=project_id, supervisor=supervisor)

        topic_title_vi = request.data.get("topic_title_vi", "").strip()
        topic_title_en = request.data.get("topic_title_en", "").strip()
        update_fields = ["status"]

        if topic_title_vi:
            project.topic_title_vi = topic_title_vi
            update_fields.append("topic_title_vi")
        if topic_title_en:
            project.topic_title_en = topic_title_en
            update_fields.append("topic_title_en")

        project.status = "TOPIC_CONFIRMED"
        project.save(update_fields=update_fields)

        try:
            NotificationService.create_notification(
                user=project.student.user,
                notification_type="general",
                title="[Đề tài ĐATN] GVHD đã xác nhận đề tài",
                message=f"GVHD {supervisor.user.get_full_name()} đã xác nhận đề tài: '{project.topic_title_vi}'. Đang trình Khoa/Bộ môn phê duyệt.",
                action_url="/student/dashboard",
                send_email=True,
            )
        except Exception as e:
            logger.warning("Could not send notification: %s", e)

        return Response({
            "message": "GV đã xác nhận đề tài thành công! Đang trình Khoa/Ban duyệt.",
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class AdminApproveTopicAPIView(APIView):
    """Trình duyệt: Khoa/Ban duyệt đề tài (Approved hoặc Yêu cầu sửa quay về Draft)"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        project_id = request.data.get("project_id")
        decision = request.data.get("decision", "APPROVED").upper()  # APPROVED or REVISION
        notes = request.data.get("notes", "").strip()

        project = get_object_or_404(GraduationProject, id=project_id)

        if decision == "APPROVED":
            project.status = "TOPIC_APPROVED"
            msg = "Khoa đã phê duyệt đề tài chính thức!"
        else:
            project.status = "TOPIC_REVISION"
            msg = f"Khoa yêu cầu chỉnh sửa đề tài: {notes or 'Vui lòng trao đổi lại với GVHD'}."

        project.save(update_fields=["status"])

        # Thông báo tới SV và GV
        try:
            for recipient in [project.student.user, project.supervisor.user]:
                NotificationService.create_notification(
                    user=recipient,
                    notification_type="general",
                    title="[Kết quả duyệt đề tài]",
                    message=msg,
                    action_url="/student/dashboard",
                    send_email=True,
                )
        except Exception as e:
            logger.warning("Could not send notification: %s", e)

        return Response({
            "message": msg,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class ExportOutlinePdfAPIView(APIView):
    """Hệ thống sinh biểu mẫu đề cương đồ án tốt nghiệp ra file PDF chuẩn UTC"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        project = get_object_or_404(
            GraduationProject.objects.select_related("student__user", "supervisor__user", "topic_category", "batch"),
            id=pk
        )

        import io
        import os
        import unicodedata
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        # Try to register a TrueType font supporting UTF-8 Vietnamese
        font_name = "Helvetica"
        possible_fonts = [
            "C:\\Windows\\Fonts\\arial.ttf",
            "C:\\Windows\\Fonts\\times.ttf",
            "C:\\Windows\\Fonts\\tahoma.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
            "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
        ]
        has_unicode_font = False
        for fpath in possible_fonts:
            if os.path.exists(fpath):
                try:
                    pdfmetrics.registerFont(TTFont("CustomFont", fpath))
                    font_name = "CustomFont"
                    has_unicode_font = True
                    break
                except Exception:
                    pass

        def clean_txt(text: str) -> str:
            if not text:
                return ""
            if has_unicode_font:
                return str(text)
            # Normalize to ASCII fallback if only Type 1 Helvetica font is available
            normalized = unicodedata.normalize('NFKD', str(text)).encode('ASCII', 'ignore').decode('ASCII')
            return normalized or str(text)

        buffer = io.BytesIO()
        p = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        # 1. Header: Bộ GD&ĐT & Quốc hiệu tiêu ngữ chuẩn biểu mẫu UTC
        p.setFont(font_name, 11)
        p.drawString(50, height - 50, clean_txt("BỘ GIÁO DỤC VÀ ĐÀO TẠO"))
        p.drawString(50, height - 66, clean_txt("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI"))
        p.drawString(50, height - 82, clean_txt("KHOA CÔNG NGHỆ THÔNG TIN"))

        p.drawString(width - 250, height - 50, clean_txt("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"))
        p.drawString(width - 210, height - 66, clean_txt("Độc lập - Tự do - Hạnh phúc"))
        p.setLineWidth(0.5)
        p.line(width - 200, height - 74, width - 100, height - 74)

        # 2. Tiêu đề Đề cương chuẩn
        p.setFont(font_name, 15)
        p.drawCentredString(width / 2, height - 120, clean_txt("ĐỀ CƯƠNG CHI TIẾT ĐỒ ÁN TỐT NGHIỆP"))

        # 3. Thông tin sinh viên & đề tài
        p.setFont(font_name, 10)
        y = height - 155

        p.drawString(50, y, clean_txt(f"1. Tên đề tài (Tiếng Việt): {project.topic_title_vi}"))
        y -= 20
        p.drawString(50, y, clean_txt(f"   Tên đề tài (Tiếng Anh): {project.topic_title_en or 'N/A'}"))
        y -= 20
        st_name = project.student.user.get_full_name() if project.student and project.student.user else "N/A"
        st_reg = project.student.registration_no if project.student else "N/A"
        st_prog = project.student.get_degree_program_display() if project.student else "Cử nhân"
        st_dept = project.student.department or "Công nghệ thông tin" if project.student else "CNTT"
        p.drawString(50, y, clean_txt(f"2. Sinh viên thực hiện: {st_name} - MSSV: {st_reg}"))
        y -= 20
        p.drawString(50, y, clean_txt(f"   Chương trình đào tạo: {st_prog} - Ngành/Lớp: {st_dept}"))
        y -= 20

        sup_name = project.supervisor.user.get_full_name() if project.supervisor and project.supervisor.user else "N/A"
        sup_title = project.supervisor.academic_title or "GV" if project.supervisor else "GV"
        p.drawString(50, y, clean_txt(f"3. Giảng viên hướng dẫn: {sup_title}. {sup_name}"))
        y -= 20
        cat_name = project.topic_category.name if project.topic_category else "Công nghệ phần mềm & Trí tuệ nhân tạo"
        batch_info = f"{project.batch.batch_name} ({project.batch.batch_code})" if project.batch else "Học kỳ tốt nghiệp"
        p.drawString(50, y, clean_txt(f"4. Hướng nghiên cứu: {cat_name}"))
        y -= 20
        p.drawString(50, y, clean_txt(f"5. Đợt thực hiện: {batch_info}"))
        y -= 28

        # 4. Mục tiêu và nội dung nghiên cứu chính
        p.drawString(50, y, clean_txt("6. Mục tiêu nghiên cứu:"))
        y -= 16
        p.drawString(65, y, clean_txt("- Khảo sát hiện trạng, phân tích nghiệp vụ và thu thập yêu cầu hệ thống."))
        y -= 16
        p.drawString(65, y, clean_txt("- Nghiên cứu các giải pháp công nghệ, thuật toán và mô hình kiến trúc phù hợp."))
        y -= 16
        p.drawString(65, y, clean_txt("- Thiết kế kiến trúc giải pháp, cơ sở dữ liệu và các giao diện người dùng."))
        y -= 24

        p.drawString(50, y, clean_txt("7. Kế hoạch và nội dung triển khai:"))
        y -= 16
        p.drawString(65, y, clean_txt("- Giai đoạn 1 (Tuần 1-4): Hoàn thành đề cương, khảo sát thực tế và phân tích yêu cầu."))
        y -= 16
        p.drawString(65, y, clean_txt("- Giai đoạn 2 (Tuần 5-10): Thiết kế hệ thống, xây dựng CSDL và lập trình các module."))
        y -= 16
        p.drawString(65, y, clean_txt("- Giai đoạn 3 (Tuần 11-14): Kiểm thử phần mềm, đánh giá hiệu năng và viết thuyết minh."))
        y -= 16
        p.drawString(65, y, clean_txt("- Giai đoạn 4 (Tuần 15): Báo cáo nghiệm thu trước Hội đồng chấm Đồ án tốt nghiệp."))
        y -= 24

        p.drawString(50, y, clean_txt("8. Dự kiến sản phẩm bàn giao:"))
        y -= 16
        p.drawString(65, y, clean_txt("- Quyển báo cáo thuyết minh Đồ án tốt nghiệp hoàn chỉnh và slide trình chiếu."))
        y -= 16
        p.drawString(65, y, clean_txt("- Mã nguồn chương trình (Source code) và phần mềm demo triển khai thực tế."))
        y -= 45

        # 5. Khối chữ ký xác nhận chuẩn
        p.drawString(80, y, clean_txt("GIẢNG VIÊN HƯỚNG DẪN"))
        p.drawString(80, y - 14, clean_txt("(Ký và ghi rõ họ tên)"))
        p.drawString(width - 220, y, clean_txt("SINH VIÊN THỰC HIỆN"))
        p.drawString(width - 220, y - 14, clean_txt("(Ký và ghi rõ họ tên)"))

        p.showPage()
        p.save()

        buffer.seek(0)
        filename = f"De_cuong_{st_reg}.pdf"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["Content-Length"] = len(buffer.getvalue())
        return response


class UploadSignedOutlineAPIView(APIView):
    """SV ký nộp bản đề cương đã ký, Khoa lưu hồ sơ (Giới hạn tối đa 5MB, chặn .exe và parse MIME type)"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, pk):
        user = request.user
        student = getattr(user, "student_profile", None)
        project = get_object_or_404(GraduationProject, id=pk)

        signed_file = request.FILES.get("signed_outline_file")
        if not signed_file:
            return Response({"detail": "Vui lòng chọn file scan đề cương đã ký."}, status=status.HTTP_400_BAD_REQUEST)

        # 1. Kiểm tra kích thước tệp tối đa 5MB (5 * 1024 * 1024)
        MAX_SIZE = 5 * 1024 * 1024
        if signed_file.size > MAX_SIZE:
            return Response(
                {"signed_outline_file": ["Kích thước tệp vượt quá giới hạn 5MB. Vui lòng tải lên file nhỏ hơn hoặc bằng 5MB."]},
                status=status.HTTP_400_BAD_REQUEST
            )

        # 2. Kiểm tra phần mở rộng tệp (chỉ cho phép .pdf)
        import os
        ext = os.path.splitext(signed_file.name)[1].lower()
        if ext != ".pdf":
            return Response(
                {"signed_outline_file": [f"Định dạng tệp '{ext}' không hợp lệ. Chỉ chấp nhận tệp định dạng PDF (.pdf). Tuyệt đối không cho phép tệp thực thi (.exe)."]},
                status=status.HTTP_400_BAD_REQUEST
            )

        # 3. Phân tích và kiểm tra MIME type từ request header / file metadata
        content_type = getattr(signed_file, "content_type", "").lower()
        blocked_mimes = [
            "application/x-msdownload", "application/x-dosexec", "application/x-executable",
            "application/x-msdos-program", "application/x-bat", "application/x-sh",
            "application/x-sharedlib", "application/octet-stream"
        ]
        if any(bad_mime in content_type for bad_mime in ["x-msdownload", "dosexec", "x-executable"]):
            return Response(
                {"signed_outline_file": ["Phát hiện tệp thực thi nguy hiểm (.exe). Hệ thống từ chối lưu trữ."]},
                status=status.HTTP_400_BAD_REQUEST
            )
        if content_type and "pdf" not in content_type and content_type not in ["application/pdf", "application/x-pdf"]:
            return Response(
                {"signed_outline_file": [f"MIME type '{content_type}' không khớp với định dạng PDF hợp lệ."]},
                status=status.HTTP_400_BAD_REQUEST
            )

        # 4. Kiểm tra Header Magic Bytes trong ruột tệp
        try:
            signed_file.seek(0)
            header = signed_file.read(512)
            signed_file.seek(0)
            # Chặn triệt để tệp nhị phân PE / Windows executable (.exe header MZ)
            if header.startswith(b"MZ"):
                return Response(
                    {"signed_outline_file": ["Phát hiện tệp thực thi Windows (.exe/MZ). Hệ thống từ chối lưu trữ vì lý do bảo mật."]},
                    status=status.HTTP_400_BAD_REQUEST
                )
            if not header.startswith(b"%PDF"):
                return Response(
                    {"signed_outline_file": ["Tệp không đúng định dạng PDF hợp lệ (Magic bytes không bắt đầu bằng %PDF)."]},
                    status=status.HTTP_400_BAD_REQUEST
                )
        except Exception:
            return Response({"signed_outline_file": ["Không thể đọc nội dung tệp tin."]}, status=status.HTTP_400_BAD_REQUEST)

        # 5. Chạy validator chuẩn hệ thống
        try:
            validate_uploaded_file(signed_file, allowed_extensions=[".pdf"], max_size_bytes=MAX_SIZE)
        except Exception as e:
            err_msg = getattr(e, "detail", str(e))
            return Response({"signed_outline_file": [str(err_msg)]}, status=status.HTTP_400_BAD_REQUEST)

        project.signed_outline_file = signed_file
        project.save(update_fields=["signed_outline_file"])

        return Response({
            "message": "Nộp bản đề cương có chữ ký thành công! Hồ sơ đã được lưu trữ.",
            "signed_outline_file_url": project.signed_outline_file.url,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# PHASE 4: THESIS ELIGIBILITY CHECK & FORCE APPROVE APIS
# ==============================================================================

class CheckThesisEligibilityAPIView(APIView):
    """Giai đoạn 4: Xét điều kiện làm đồ án (CPA >= 2.0, nợ tín chỉ, môn học/tự chọn)"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        project = get_object_or_404(GraduationProject, id=pk)
        
        # Hỗ trợ nhận thông tin nợ tín chỉ và dữ liệu môn học từ request body
        debt_credits = request.data.get("debt_credits")
        if debt_credits is None:
            debt_credits = request.data.get("credits_debt") or request.data.get("no_tin_chi")
        if debt_credits is not None:
            try:
                debt_credits = int(debt_credits)
            except (ValueError, TypeError):
                debt_credits = None

        courses_data = request.data.get("courses") or request.data.get("mon_hoc") or request.data.get("transcript") or request.data.get("electives")

        ok, msg = AcademicClearanceService.check_thesis_start_eligibility(
            project, debt_credits=debt_credits, courses_data=courses_data
        )
        return Response({
            "success": True,
            "is_eligible": ok,
            "message": msg,
            "status": project.status,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


class ForceApproveThesisAPIView(APIView):
    """Giai đoạn 4: Nút quyết định Force Approve của Khoa (vẫn cho làm ĐA dù chưa đủ tiêu chí)"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        project = get_object_or_404(GraduationProject, id=pk)
        force = request.data.get("force", True)
        if isinstance(force, str):
            force = force.lower() in ["true", "1"]

        if force:
            project.is_force_approved = True
            project.status = "IN_PROGRESS"
            msg = f"Khoa đã duyệt đặc cách (Force Approve) cho SV {project.student.registration_no} tiếp tục làm đồ án!"
        else:
            project.is_force_approved = False
            project.status = "DISQUALIFIED"
            msg = f"Không duyệt đặc cách. SV {project.student.registration_no} bị loại khỏi đợt đồ án."

        project.save(update_fields=["is_force_approved", "status"])

        return Response({
            "success": True,
            "message": msg,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# PHASE 5: TASK DELIVERABLES & SUPERVISOR REVIEW APIS
# ==============================================================================

class StudentTaskDeliverableSubmitAPIView(APIView):
    """SV submit kết quả task (file báo cáo, link demo/commit, ghi chú)"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request, pk):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Chỉ dành cho sinh viên."}, status=status.HTTP_403_FORBIDDEN)

        task = get_object_or_404(SupervisionTask, id=pk, project__student=student)
        deliverable_file = request.FILES.get("deliverable_file")
        deliverable_url = request.data.get("deliverable_url", "").strip()
        student_notes = request.data.get("student_notes", "").strip()

        if deliverable_file:
            try:
                validate_uploaded_file(
                    deliverable_file,
                    allowed_extensions=[".pdf", ".zip", ".rar", ".docx", ".xlsx", ".pptx"],
                    max_size_bytes=25 * 1024 * 1024
                )
                task.deliverable_file = deliverable_file
            except Exception as e:
                err_msg = getattr(e, "detail", str(e))
                return Response({"deliverable_file": [str(err_msg)]}, status=status.HTTP_400_BAD_REQUEST)

        if deliverable_url:
            task.deliverable_url = deliverable_url
        if student_notes:
            task.student_notes = student_notes

        task.status = "IN_PROGRESS"
        task.review_verdict = "PENDING"
        task.save()

        # Thông báo tới GVHD
        try:
            NotificationService.create_notification(
                user=task.project.supervisor.user,
                notification_type="general",
                title="[Nộp kết quả nhiệm vụ]",
                message=f"SV {student.user.get_full_name()} đã nộp kết quả cho nhiệm vụ: '{task.title}'. Đang chờ Thầy/Cô đánh giá.",
                action_url="/supervisor/dashboard",
                send_email=True,
            )
        except Exception as e:
            logger.warning("Could not send notification: %s", e)

        return Response({
            "message": "Nộp kết quả nhiệm vụ thành công! Đang chờ GVHD đánh giá.",
            "task": SupervisionTaskSerializer(task).data
        }, status=status.HTTP_200_OK)


class SupervisorReviewTaskAPIView(APIView):
    """GV review kết quả task: Nút quyết định Đạt? (ACCEPTED: Đạt / REVISION_REQUIRED: Yêu cầu sửa)"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        user = request.user
        supervisor = getattr(user, "supervisor_profile", None)
        if not supervisor:
            return Response({"detail": "Chỉ dành cho Giảng viên hướng dẫn."}, status=status.HTTP_403_FORBIDDEN)

        task = get_object_or_404(SupervisionTask, id=pk, project__supervisor=supervisor)
        verdict = request.data.get("verdict", "ACCEPTED").upper()  # ACCEPTED or REVISION_REQUIRED
        notes = request.data.get("notes", "").strip()

        if verdict == "ACCEPTED":
            task.review_verdict = "ACCEPTED"
            task.status = "COMPLETED"
            task.is_completed = True
            task.completed_at = timezone.now()
            msg = f"GVHD đã đánh giá ĐẠT cho nhiệm vụ '{task.title}'!"
        else:
            task.review_verdict = "REVISION_REQUIRED"
            task.status = "IN_PROGRESS"
            task.is_completed = False
            task.completed_at = None
            msg = f"GVHD yêu cầu làm lại nhiệm vụ '{task.title}': {notes or 'Cần hoàn thiện thêm'}."

        task.supervisor_review_notes = notes
        task.reviewed_at = timezone.now()
        task.save()

        # Thông báo tới SV
        try:
            NotificationService.create_notification(
                user=task.project.student.user,
                notification_type="general",
                title="[Đánh giá nhiệm vụ từ GVHD]",
                message=msg,
                action_url="/student/dashboard",
                send_email=True,
            )
        except Exception as e:
            logger.warning("Could not send notification: %s", e)

        return Response({
            "message": msg,
            "task": SupervisionTaskSerializer(task).data
        }, status=status.HTTP_200_OK)


# ==============================================================================
# PHASE 6: EXPORT MINUTES PDF & BATCH FINAL GRADES EXCEL APIS
# ==============================================================================

class CouncilMinutesPdfExportAPIView(APIView):
    """Biên bản: Xuất biên bản họp bảo vệ PDF theo chuẩn UTC (1CT-2TK-2UV)"""
    permission_classes = [IsAuthenticated]

    def get(self, request, council_id):
        council = get_object_or_404(
            DefenseCouncil.objects.select_related("batch"),
            id=council_id
        )

        import io
        import os
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        font_name = "Helvetica"
        if os.path.exists("C:\\Windows\\Fonts\\arial.ttf"):
            try:
                pdfmetrics.registerFont(TTFont("Arial", "C:\\Windows\\Fonts\\arial.ttf"))
                font_name = "Arial"
            except Exception:
                pass

        buffer = io.BytesIO()
        p = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        # Header
        p.setFont(font_name, 11)
        p.drawString(50, height - 50, "BỘ GIÁO DỤC VÀ ĐÀO TẠO")
        p.drawString(50, height - 68, "TRƯỜNG ĐH GIAO THÔNG VẬN TẢI")
        p.drawString(width - 240, height - 50, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
        p.drawString(width - 200, height - 68, "Độc lập - Tự do - Hạnh phúc")

        p.setFont(font_name, 15)
        p.drawCentredString(width / 2, height - 110, f"BIÊN BẢN HỌP HỘI ĐỒNG BẢO VỆ ĐỒ ÁN TỐT NGHIỆP #{council.council_number}")

        p.setFont(font_name, 10)
        y = height - 140
        session_time_disp = council.get_session_time_display()
        p.drawString(50, y, f"Hội đồng: {council.council_name} | Phòng: {council.defense_room or 'TBA'} | Ngày: {council.session_date or 'TBA'} ({session_time_disp})")
        y -= 25

        p.drawString(50, y, "I. THÀNH VIÊN HỘI ĐỒNG (Cơ cấu chuẩn 1CT - 2TK - 2UV):")
        y -= 18
        members = list(council.members.select_related("user").all())
        for idx, m in enumerate(members, 1):
            p.drawString(70, y, f"{idx}. {m.get_role_display()}: {m.user.get_full_name()} ({m.user.email})")
            y -= 16

        y -= 15
        p.drawString(50, y, "II. KẾT QUẢ ĐÁNH GIÁ CỦA CÁC SINH VIÊN:")
        y -= 20

        # Projects and grades
        projects = council.projects.select_related("student__user", "final_grade_summary").all()
        p.drawString(50, y, "STT | MSSV | Họ và tên | Điểm GVHD | Điểm GVPB | Điểm TB HĐ | Tổng kết (Hệ 10 / Chữ)")
        y -= 15

        for idx, proj in enumerate(projects, 1):
            summary = getattr(proj, "final_grade_summary", None)
            sup_s = f"{proj.supervisor_score:.1f}" if proj.supervisor_score is not None else "--"
            rev_s = f"{proj.reviewer_score:.1f}" if proj.reviewer_score is not None else "--"
            cou_s = f"{summary.council_avg_score:.2f}" if summary and summary.council_avg_score is not None else "--"
            fin_s = f"{summary.final_score_10:.2f} ({summary.final_letter_grade})" if summary and summary.final_score_10 is not None else "--"
            s_name = proj.student.user.get_full_name()
            p.drawString(50, y, f"{idx}. {proj.student.registration_no} - {s_name} | {sup_s}đ | {rev_s}đ | {cou_s}đ | {fin_s}")
            y -= 16
            if y < 100:
                p.showPage()
                y = height - 50

        y -= 40
        p.drawString(70, y, "CHỦ TỊCH HỘI ĐỒNG")
        p.drawString(width - 220, y, "THƯ KÝ HỘI ĐỒNG")
        p.drawString(70, y - 15, "(Ký và ghi rõ họ tên)")
        p.drawString(width - 220, y - 15, "(Ký và ghi rõ họ tên)")

        p.showPage()
        p.save()

        buffer.seek(0)
        filename = f"Bien_ban_HD_{council.council_number}.pdf"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class BatchFinalGradesExcelExportAPIView(APIView):
    """Xuất cuối kỳ: Xuất dữ liệu cuối kỳ toàn bộ đợt đồ án ra Excel (ký, lưu)"""
    permission_classes = [IsAuthenticated]

    def get(self, request, batch_id):
        batch = get_object_or_404(AcademicBatch, id=batch_id)
        from openpyxl import Workbook
        import io

        wb = Workbook()
        ws = wb.active
        ws.title = "Bảng điểm tốt nghiệp"

        ws.append(["TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI - KHOA CÔNG NGHỆ THÔNG TIN"])
        ws.append([f"BẢNG TỔNG HỢP KẾT QUẢ ĐỒ ÁN TỐT NGHIỆP: {batch.batch_name}"])
        ws.append([])
        ws.append([
            "STT", "MSSV", "Họ và tên", "Chương trình", "Tên đề tài tiếng Việt", "GVHD",
            "Hội đồng", "Điểm GVHD", "Điểm GVPB", "Điểm TB HĐ", "Điểm Tổng (10)",
            "Điểm Hệ 4", "Điểm chữ", "Xếp loại", "Kết quả"
        ])

        projects = GraduationProject.objects.filter(batch=batch).select_related(
            "student__user", "supervisor__user", "council", "final_grade_summary"
        ).order_by("student__registration_no")

        for idx, p in enumerate(projects, 1):
            s = getattr(p, "final_grade_summary", None)
            ws.append([
                idx,
                p.student.registration_no,
                p.student.user.get_full_name(),
                p.student.get_degree_program_display(),
                p.topic_title_vi,
                p.supervisor.user.get_full_name(),
                p.council.council_name if p.council else "Chưa gán",
                p.supervisor_score or "",
                p.reviewer_score or "",
                s.council_avg_score if s else "",
                s.final_score_10 if s else "",
                s.final_score_4 if s else "",
                s.final_letter_grade if s else "",
                s.classification if s else "",
                "Đạt" if (s and s.is_passed) else ("Bảo lưu" if p.status == "DEFERRED" else "Không đạt")
            ])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"Bang_diem_tot_nghiep_{batch.batch_code}.xlsx"
        response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


# ==============================================================================
# DEFERRED THESIS (NHÁNH BẢO LƯU) APIS
# ==============================================================================

class StudentDeferralRequestAPIView(APIView):
    """Sinh viên nộp đơn xin bảo lưu đồ án khi không đủ điều kiện bảo vệ / học vụ"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Chỉ dành cho sinh viên."}, status=status.HTTP_403_FORBIDDEN)

        requests = ThesisDeferralRequest.objects.filter(student=student).order_by("-submitted_at")
        return Response(ThesisDeferralRequestSerializer(requests, many=True).data, status=status.HTTP_200_OK)

    def post(self, request):
        user = request.user
        student = getattr(user, "student_profile", None)
        if not student:
            return Response({"detail": "Chỉ dành cho sinh viên."}, status=status.HTTP_403_FORBIDDEN)

        project = GraduationProject.objects.filter(student=student).first()
        if not project:
            batch = getattr(student, "academic_batch", None) or AcademicBatch.objects.filter(is_active=True).first()
            if batch:
                project = GraduationProject.objects.create(
                    student=student,
                    batch=batch,
                    status="DISQUALIFIED"
                )
            else:
                return Response({"detail": "Chưa có đồ án hoặc đợt hợp lệ để xin bảo lưu."}, status=status.HTTP_400_BAD_REQUEST)

        reason = (request.data.get("reason", "") or request.data.get("reason_details", "")).strip()
        evidence_file = request.FILES.get("evidence_file")

        if not reason:
            return Response({"detail": "Vui lòng nhập lý do xin bảo lưu đồ án."}, status=status.HTTP_400_BAD_REQUEST)

        if evidence_file:
            try:
                validate_uploaded_file(evidence_file, allowed_extensions=[".pdf", ".zip", ".docx", ".png", ".jpg"], max_size_bytes=25 * 1024 * 1024)
            except Exception as e:
                err_msg = getattr(e, "detail", str(e))
                return Response({"evidence_file": [str(err_msg)]}, status=status.HTTP_400_BAD_REQUEST)

        deferral = ThesisDeferralRequest.objects.create(
            project=project,
            student=student,
            reason=reason,
            evidence_file=evidence_file,
            status="PENDING",
        )

        return Response({
            "success": True,
            "message": "Nộp đơn xin bảo lưu đồ án thành công! Đang chờ Khoa phê duyệt.",
            "deferral": ThesisDeferralRequestSerializer(deferral).data
        }, status=status.HTTP_201_CREATED)


class AdminReviewDeferralRequestAPIView(APIView):
    """Khoa duyệt đơn bảo lưu: Có -> status='DEFERRED', lưu hồ sơ; Không -> loại khỏi đợt"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        deferral = get_object_or_404(ThesisDeferralRequest, id=pk)
        decision = request.data.get("decision", "APPROVED").upper()  # APPROVED or REJECTED
        admin_notes = request.data.get("admin_notes", "").strip()

        with transaction.atomic():
            deferral.status = decision
            deferral.admin_notes = admin_notes
            deferral.reviewed_at = timezone.now()
            deferral.reviewed_by = request.user
            deferral.save()

            project = deferral.project
            if decision == "APPROVED":
                project.status = "DEFERRED"
                msg = f"Khoa đã duyệt đơn bảo lưu đồ án cho sinh viên {project.student.registration_no}."
            else:
                project.status = "DISQUALIFIED"
                msg = f"Khoa không duyệt đơn bảo lưu. Sinh viên {project.student.registration_no} bị loại khỏi đợt."

            project.save(update_fields=["status"])

        # Thông báo tới SV
        try:
            NotificationService.create_notification(
                user=deferral.student.user,
                notification_type="general",
                title="[Kết quả duyệt đơn bảo lưu đồ án]",
                message=msg,
                action_url="/student/dashboard",
                send_email=True,
            )
        except Exception as e:
            logger.warning("Could not send notification: %s", e)

        return Response({
            "message": msg,
            "deferral": ThesisDeferralRequestSerializer(deferral).data,
            "project": GraduationProjectDetailSerializer(project).data
        }, status=status.HTTP_200_OK)

